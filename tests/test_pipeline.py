"""Tests for the rules in specification_v1.md."""

from __future__ import annotations

import shutil
from datetime import date, datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

from datatransform.coerce import to_date, to_number, to_text
from datatransform.crosschecks import (
    hypotheses_from,
    match_record,
    run_rules,
    split_label,
)
from datatransform.extract import extract_sheet, last_non_empty_row
from datatransform.markers import attributes_for, block_indices, parse_marker, read_markers
from datatransform.model import (
    Confidence,
    ExtractionError,
    FieldType,
    Orientation,
    Rule,
    parse_reference,
)
from datatransform.nomenclature import match_sheet, norm, read_nomenclature, sheet_sort_key
from datatransform.runner import run
from datatransform.specs import step2_for
from datatransform.transform import apply_step2
from datatransform.writer import ANCHOR_PREFIX

SOURCE = Path(__file__).resolve().parents[1] / "Intake_v1.xlsx"
ROW_WISE = "01. History"
TRANSPOSED = "01. History_Transposed"

# Year is read as text, measures as float — spec §8.4.
# 2025 appears twice: the full year and its first nine months — spec §9.2 S10.
EXPECTED = [
    ("2021", 15900.0, 14930.0),
    ("2022", 16740.0, 10030.0),
    ("2023", 17520.0, 12660.0),
    ("2024", 18390.0, 11700.0),
    ("2025", 19200.0, 10250.0),
    ("2025 9 months", 14400.0, 7100.0),
]
TOTALS = {"Premium": 102150.0, "Incurred Losses": 66670.0}
DISTINCT = ["2019", "2020", "2021", "2022", "2023", "2024"]


@pytest.fixture(scope="module")
def books():
    return (load_workbook(SOURCE, data_only=True), load_workbook(SOURCE, data_only=False))


@pytest.fixture(scope="module")
def nomenclature(books):
    return read_nomenclature(books[1])


def _block(books, nomenclature, sheet):
    values, formulas = books
    markers = read_markers(formulas[sheet], last_non_empty_row(values[sheet]))
    blocks = extract_sheet(values[sheet], formulas[sheet], nomenclature, markers)
    assert len(blocks) == 1
    return blocks[0]


# ───────────────────────────────────────────────────────── §4 marker grammar

@pytest.mark.parametrize("text,name,value,index,hyp", [
    ("Header_1", "Header", None, 1, False),
    ("Info_1 = L", "Info", "L", 1, False),
    ("Info_1 = 16", "Info", "16", 1, False),
    ("Transpose_1", "Transpose", None, 1, False),
    ("Currency = USD", "Currency", "USD", None, False),
    ("H_Year basis = UW", "Year basis", "UW", None, True),
    ("H_PF transfer = with clean cut", "PF transfer", "with clean cut", None, True),
    ("Currency_2 = EUR", "Currency", "EUR", 2, False),
])
def test_marker_grammar(text, name, value, index, hyp):
    m = parse_marker(text, 1)
    assert (m.name, m.value, m.index, m.is_hypothesis) == (name, value, index, hyp)


def test_block_declared_only_by_header(books, nomenclature):
    values, formulas = books
    markers = read_markers(formulas[ROW_WISE])
    assert block_indices(markers) == [1]


def test_block_override_beats_sheet_wide_attribute():
    markers = [parse_marker("Currency = USD", 1), parse_marker("Currency_2 = EUR", 2)]
    assert attributes_for(markers, 1)["Currency"].value == "USD"
    assert attributes_for(markers, 2)["Currency"].value == "EUR"


def test_no_header_marker_extracts_nothing(books, nomenclature):
    values, formulas = books
    blocks = extract_sheet(values[ROW_WISE], formulas[ROW_WISE], nomenclature, markers=[])
    assert blocks == []


# ───────────────────────────────────────────────────── §2.2, §8 name matching

def test_numeric_prefix_matching_survives_misspelling():
    assert sheet_sort_key("10. Triangles")[0] == sheet_sort_key("10. Triangels")[0]
    assert match_sheet("10. Triangles", ["10. Triangels"]) == "10. Triangels"


def test_norm_strips_nbsp_and_quote_style():
    assert norm("  Premium  ") == "Premium"
    assert norm("„01. History“") == '"01. History"'


# ───────────────────────────────────────────────────────── §3 nomenclature

def test_nomenclature_declares_dataset_01(nomenclature):
    ds = nomenclature.datasets[ROW_WISE]
    assert ds.key == "01 History"
    assert ds.headers == ("Year", "Premium", "Incurred Losses")
    assert ds.key_field == "Year"
    assert "PF transfer" in ds.attributes


def test_vocabulary_is_enforced_only_where_defined(nomenclature):
    assert nomenclature.validate_value("Year basis", "UW")
    assert not nomenclature.validate_value("Year basis", "Underwriting")
    assert nomenclature.validate_value("Currency", "USD")
    assert not nomenclature.validate_value("Currency", "Dollars")
    assert nomenclature.validate_value("As at", "31.12.2025")     # no vocabulary entry


# ───────────────────────────────────────────────── §7, §8 resolution & selection

def test_row_wise_resolution(books, nomenclature):
    b = _block(books, nomenclature, ROW_WISE)
    assert b.orientation is Orientation.ROW_WISE
    assert b.header_ref == "8" and b.info_ref == "L"
    assert b.address_map == {"Year": "B", "Premium": "D", "Incurred Losses": "G"}
    assert b.provenance_label == "Source row"


def test_transposed_resolution(books, nomenclature):
    b = _block(books, nomenclature, TRANSPOSED)
    assert b.orientation is Orientation.TRANSPOSED
    assert b.header_ref == "C" and b.info_ref == "16"
    assert b.address_map == {"Year": "8", "Premium": "10", "Incurred Losses": "13"}
    assert b.provenance_label == "Source column"


@pytest.mark.parametrize("sheet", [ROW_WISE, TRANSPOSED])
def test_records_and_exclusion(books, nomenclature, sheet):
    b = _block(books, nomenclature, sheet)
    assert (b.candidates, len(b.records), b.excluded) == (7, 6, 1)
    got = [(r.values["Year"], r.values["Premium"], r.values["Incurred Losses"])
           for r in b.records]
    assert got == EXPECTED


def test_both_orientations_are_equivalent(books, nomenclature):
    a = _block(books, nomenclature, ROW_WISE)
    b = _block(books, nomenclature, TRANSPOSED)
    assert [r.values for r in a.records] == [r.values for r in b.records]
    assert a.totals() == b.totals() == TOTALS
    assert a.provenance_label != b.provenance_label


def test_excluded_total_ties_to_extracted_total(books, nomenclature):
    b = _block(books, nomenclature, ROW_WISE)
    assert len(b.excluded_records) == 1
    excluded = b.excluded_records[0].values
    assert excluded["Premium"] == b.totals()["Premium"]
    assert excluded["Incurred Losses"] == b.totals()["Incurred Losses"]


def test_undeclared_columns_are_logged_not_dropped(books, nomenclature):
    assert _block(books, nomenclature, ROW_WISE).unextracted == ["C", "E", "F"]
    assert _block(books, nomenclature, TRANSPOSED).unextracted == ["9", "11", "12"]


def test_only_declared_addresses_are_read(books, nomenclature):
    b = _block(books, nomenclature, ROW_WISE)
    assert len(b.records) * len(b.dataset.headers) == 18
    assert all(set(r.values) == set(b.dataset.headers) for r in b.records)


# ───────────────────────────────────────────────────────── §5 confidence

def test_confidence_is_derived_and_worst_wins(books, nomenclature):
    b = _block(books, nomenclature, ROW_WISE)
    assert {h.attribute for h in b.hypotheses} == {
        "Year basis", "Loss basis", "PF transfer", "Period overlap 2025",
    }
    assert all(h.confidence is Confidence.ASSUMED for h in b.hypotheses)
    assert b.confidence is Confidence.ASSUMED


def test_open_ranks_worse_than_assumed():
    assert Confidence.OPEN.rank > Confidence.ASSUMED.rank > Confidence.CONFIRMED.rank
    assert Confidence.worst([Confidence.CONFIRMED, Confidence.OPEN]) is Confidence.OPEN
    assert Confidence.worst([]) is Confidence.CONFIRMED


def test_missing_attribute_raises_open_hypothesis(books, nomenclature, monkeypatch):
    values, formulas = books
    ds = nomenclature.datasets[ROW_WISE]
    markers = [m for m in read_markers(formulas[ROW_WISE], last_non_empty_row(values[ROW_WISE]))
               if m.name != "PF transfer"]
    blocks = extract_sheet(values[ROW_WISE], formulas[ROW_WISE], nomenclature, markers)
    opens = [h for h in blocks[0].hypotheses if h.confidence is Confidence.OPEN]
    assert [h.attribute for h in opens] == ["PF transfer"]
    assert blocks[0].confidence is Confidence.OPEN


def test_row_level_confidence_only_when_transposed(books, nomenclature):
    assert all(r.confidence is None for r in _block(books, nomenclature, ROW_WISE).records)
    assert all(r.confidence is Confidence.ASSUMED
               for r in _block(books, nomenclature, TRANSPOSED).records)


# ───────────────────────────────────────────────────────── §9.2 step 2

def test_step2_sorts_reorders_and_calculates(books, nomenclature):
    b = _block(books, nomenclature, ROW_WISE)
    result = apply_step2(b, step2_for(b.dataset.key))
    assert [r.values["Year"] for r in result.records] == [
        "2021", "2022", "2023", "2024", "2025", "2025 9 months",
    ]
    assert result.columns == ("Year", "Premium", "Incurred Losses", "Loss Ratio %")
    ratios = result.computed["Loss Ratio %"]
    assert ratios[0] == pytest.approx(14930 / 15900)
    assert ratios[-2] == pytest.approx(10250 / 19200)     # 2025 full year
    assert ratios[-1] == pytest.approx(7100 / 14400)      # 2025 first nine months


def test_step2_is_value_preserving(books, nomenclature):
    b = _block(books, nomenclature, ROW_WISE)
    result = apply_step2(b, step2_for(b.dataset.key))
    assert result.totals() == b.totals()


def test_step2_guard_catches_a_value_changing_bug(books, nomenclature, monkeypatch):
    """Sorting cannot move a total; if it does, that is a bug — spec §10 C3."""
    from datatransform import transform as tmod

    b = _block(books, nomenclature, ROW_WISE)
    monkeypatch.setattr(tmod, "sorted", lambda seq, **kw: list(seq)[:-1], raising=False)
    with pytest.raises(ExtractionError, match="value-preserving"):
        apply_step2(b, step2_for(b.dataset.key))


# ───────────────────────────────────────────────────────── §9 end to end

@pytest.fixture
def workspace(tmp_path):
    src = tmp_path / "Intake.xlsx"
    shutil.copy2(SOURCE, src)
    return src


def test_run_writes_blocks_and_leaves_source_untouched(workspace, tmp_path):
    before = workspace.read_bytes()
    out = tmp_path / "out.xlsx"
    report = run(workspace, out, tmp_path / "logs")

    assert report.ok
    assert workspace.read_bytes() == before, "the source workbook must never be modified"
    assert out.exists()

    wb = load_workbook(out, data_only=True)
    anchors = [
        c.value for row in wb[ROW_WISE].iter_rows() for c in row
        if isinstance(c.value, str) and c.value.startswith(ANCHOR_PREFIX)
    ]
    assert anchors == ["⟦DT:01 History:STEP1:v1⟧", "⟦DT:01 History:STEP2:v1⟧",
                       "⟦DT:per risk sections:LOSSSHARE:v1⟧"]


def test_generated_blocks_start_below_the_last_original_row(workspace, tmp_path):
    original_last = last_non_empty_row(load_workbook(workspace, data_only=True)[ROW_WISE])
    out = tmp_path / "out.xlsx"
    run(workspace, out, tmp_path / "logs")

    ws = load_workbook(out, data_only=True)[ROW_WISE]
    anchor = next(c.row for row in ws.iter_rows() for c in row
                  if isinstance(c.value, str) and c.value.startswith(ANCHOR_PREFIX))
    assert anchor == original_last + 1 + 3          # three blank rows — spec §9.1 O2


def test_rerun_replaces_rather_than_stacks(workspace, tmp_path):
    """Processing a workbook that already carries generated blocks must replace them."""
    out = tmp_path / "out.xlsx"
    run(workspace, out, tmp_path / "logs")
    first_extent = load_workbook(out, data_only=True)[ROW_WISE].max_row

    already_processed = tmp_path / "again_in.xlsx"
    shutil.copy2(out, already_processed)
    again = tmp_path / "again_out.xlsx"
    run(already_processed, again, tmp_path / "logs")

    ws = load_workbook(again, data_only=True)[ROW_WISE]
    anchors = [c.value for row in ws.iter_rows() for c in row
               if isinstance(c.value, str) and c.value.startswith(ANCHOR_PREFIX)]
    assert len(anchors) == 3, "a re-run must replace the previous blocks, not append to them"
    assert ws.max_row == first_extent


def test_output_may_not_overwrite_the_source(workspace, tmp_path):
    with pytest.raises(ExtractionError, match="never modified"):
        run(workspace, workspace, tmp_path / "logs")


def test_formulas_carry_cached_values(workspace, tmp_path):
    out = tmp_path / "out.xlsx"
    run(workspace, out, tmp_path / "logs")
    values = load_workbook(out, data_only=True)[ROW_WISE]
    formulas = load_workbook(out, data_only=False)[ROW_WISE]

    checked = 0
    for row in formulas.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cached = values[cell.coordinate].value
                assert cached is not None, f"{cell.coordinate} {cell.value} has no cached value"
                checked += 1
    assert checked >= 20


def test_control_checks_all_read_ok(workspace, tmp_path):
    out = tmp_path / "out.xlsx"
    run(workspace, out, tmp_path / "logs")
    wb = load_workbook(out, data_only=True)
    for sheet in (ROW_WISE, TRANSPOSED):
        ws = wb[sheet]
        checks = [c.value for row in ws.iter_rows() for c in row if c.value == "MISMATCH"]
        assert not checks
        oks = [c.value for row in ws.iter_rows() for c in row if c.value == "OK"]
        assert len(oks) == 4                        # two measures × step 1 and step 2


def test_logs_are_written(workspace, tmp_path):
    logs = tmp_path / "logs"
    run(workspace, tmp_path / "out.xlsx", logs)
    assert len(list(logs.glob("*_debug.log"))) == 1
    process = next(logs.glob("*_process.log")).read_text(encoding="utf-8")
    assert "sha256" in process
    assert "Fields resolved by label" in process
    assert "Dataset confidence: Assumed" in process


# ───────────────────────────────────────────────── §7.2, §8 fatal conditions

def test_uncalculated_selector_is_fatal(workspace, tmp_path, nomenclature):
    wb = load_workbook(workspace, data_only=False)
    values = load_workbook(workspace, data_only=True)
    values[ROW_WISE]["L10"] = None                  # cache gone, formula still present
    markers = read_markers(wb[ROW_WISE], last_non_empty_row(values[ROW_WISE]))
    with pytest.raises(ExtractionError, match="workbook not calculated"):
        extract_sheet(values[ROW_WISE], wb[ROW_WISE], nomenclature, markers)


def test_broken_selector_integrity_is_fatal(workspace, nomenclature):
    wb = load_workbook(workspace, data_only=False)
    values = load_workbook(workspace, data_only=True)
    values[ROW_WISE]["L10"] = 99                    # says row 99, sits on row 10
    markers = read_markers(wb[ROW_WISE], last_non_empty_row(values[ROW_WISE]))
    with pytest.raises(ExtractionError, match="manipulated"):
        extract_sheet(values[ROW_WISE], wb[ROW_WISE], nomenclature, markers)


def test_error_cell_is_fatal_not_coerced(workspace, nomenclature):
    wb = load_workbook(workspace, data_only=False)
    values = load_workbook(workspace, data_only=True)
    values[ROW_WISE]["D10"] = "#REF!"
    markers = read_markers(wb[ROW_WISE], last_non_empty_row(values[ROW_WISE]))
    with pytest.raises(ExtractionError, match="#REF!"):
        extract_sheet(values[ROW_WISE], wb[ROW_WISE], nomenclature, markers)


def test_missing_declared_label_is_fatal(workspace, nomenclature):
    wb = load_workbook(workspace, data_only=False)
    values = load_workbook(workspace, data_only=True)
    values[ROW_WISE]["D8"] = None                   # remove 'Premium' from extraction row
    markers = read_markers(wb[ROW_WISE], last_non_empty_row(values[ROW_WISE]))
    with pytest.raises(ExtractionError, match="absent from the extraction"):
        extract_sheet(values[ROW_WISE], wb[ROW_WISE], nomenclature, markers)


def test_duplicate_label_is_fatal(workspace, nomenclature):
    wb = load_workbook(workspace, data_only=False)
    values = load_workbook(workspace, data_only=True)
    values[ROW_WISE]["E8"] = "Premium"              # a second 'Premium'
    markers = read_markers(wb[ROW_WISE], last_non_empty_row(values[ROW_WISE]))
    with pytest.raises(ExtractionError, match="ambiguous"):
        extract_sheet(values[ROW_WISE], wb[ROW_WISE], nomenclature, markers)


def test_value_outside_vocabulary_is_fatal(workspace, nomenclature):
    wb = load_workbook(workspace, data_only=False)
    values = load_workbook(workspace, data_only=True)
    wb[ROW_WISE]["A4"] = "H_Year basis = Underwriting"   # markers are read from the formula handle
    markers = read_markers(wb[ROW_WISE], last_non_empty_row(values[ROW_WISE]))
    with pytest.raises(ExtractionError, match="not in its vocabulary"):
        extract_sheet(values[ROW_WISE], wb[ROW_WISE], nomenclature, markers)


# ───────────────────────────────────────────────────────── §8.4 typing

def test_declared_types_are_applied(books, nomenclature):
    b = _block(books, nomenclature, ROW_WISE)
    assert b.field_types["Year"] is FieldType.TEXT
    assert b.field_types["Premium"] is FieldType.NUMBER
    assert b.field_types["Incurred Losses"] is FieldType.NUMBER
    for r in b.records:
        assert isinstance(r.values["Year"], str)
        assert isinstance(r.values["Premium"], float)
        assert isinstance(r.values["Incurred Losses"], float)


def test_only_numeric_fields_are_summed(books, nomenclature):
    b = _block(books, nomenclature, ROW_WISE)
    assert b.numeric_fields == ("Premium", "Incurred Losses")
    assert "Year" not in b.totals()


def test_year_is_written_as_text_not_a_number(books, nomenclature):
    b = _block(books, nomenclature, ROW_WISE)
    assert b.number_format("Year") == "@"
    assert b.number_format("Premium") == "#,##0"


@pytest.mark.parametrize("raw,expected", [
    (2021, "2021"),
    (2021.0, "2021"),
    ("2021", "2021"),
    ("  2021 ", "2021"),
    ("2021/22", "2021/22"),
    (datetime(2021, 12, 31), "2021"),
    (None, None),
    ("", None),
])
def test_text_coercion(raw, expected):
    assert to_text(raw, "Year", "B9", []) == expected


@pytest.mark.parametrize("raw,expected", [
    (15900, 15900.0),
    (15900.5, 15900.5),
    ("15900", 15900.0),
    ("15 900", 15900.0),
    ("15'900", 15900.0),
    ("15,900.50", 15900.5),      # English: rightmost separator is the decimal point
    ("15.900,50", 15900.5),      # German: likewise
    ("1.234.567", 1234567.0),    # repeated separator can only be grouping
    ("12.5", 12.5),              # two trailing digits is a decimal point
    ("1.2345", 1.2345),          # four trailing digits likewise
    ("(500)", -500.0),           # accounting negative
    ("-500", -500.0),
    ("15900 USD", 15900.0),
    (None, None),
    ("", None),
])
def test_number_coercion(raw, expected):
    assert to_number(raw, "Premium", "D9", []) == expected


@pytest.mark.parametrize("raw", ["1.234", "1,234"])
def test_ambiguous_separator_is_fatal(raw):
    """1.234 is 1234 in German and 1.234 in English — never guess."""
    with pytest.raises(ExtractionError, match="ambiguous"):
        to_number(raw, "Premium", "D9", [])


def test_unparseable_number_is_fatal():
    with pytest.raises(ExtractionError, match="cannot be read as a number"):
        to_number("n/a", "Premium", "D9", [])


def test_empty_is_missing_not_zero():
    """Zero is a claim about the data; missing is an absence."""
    assert to_number(None, "Premium", "D9", []) is None
    assert to_number("   ", "Premium", "D9", []) is None
    assert to_number(0, "Premium", "D9", []) == 0.0


def test_coercions_are_recorded_for_reporting():
    log = []
    to_number("15 900", "Premium", "D9", log)
    to_text(2021, "Year", "B9", log)
    assert [c.field for c in log] == ["Premium", "Year"]
    assert log[0].before == "15 900" and log[0].after == 15900.0


def test_source_workbook_needs_no_coercion(books, nomenclature):
    """The reference workbook is already properly typed apart from Year."""
    b = _block(books, nomenclature, ROW_WISE)
    assert {c.field for c in b.coercions} == {"Year"}


def test_sort_is_numeric_aware_despite_text_years(books, nomenclature):
    b = _block(books, nomenclature, ROW_WISE)
    b.records[0].values["Year"] = "999"          # a plain string sort would put this last
    result = apply_step2(b, step2_for(b.dataset.key))
    assert [r.values["Year"] for r in result.records][0] == "999"


def test_written_year_cell_is_text_formatted(workspace, tmp_path):
    out = tmp_path / "out.xlsx"
    run(workspace, out, tmp_path / "logs")
    ws = load_workbook(out, data_only=True)[ROW_WISE]
    year_cells = [c for row in ws.iter_rows() for c in row if c.value == "2021"]
    assert year_cells, "the extracted year should be written as text"
    assert all(c.number_format == "@" for c in year_cells)


def test_routine_and_notable_coercions_are_distinguished():
    log = []
    to_text(2021, "Year", "B9", log)          # declared type applied to a clean cell
    to_number("15 900", "Premium", "D9", log)  # a number that arrived as text
    assert [c.routine for c in log] == [True, False]


def test_text_numbers_in_source_are_read_and_reported(workspace, tmp_path, nomenclature):
    """A premium formatted as text must still be read, and the conversion reported."""
    from openpyxl import load_workbook as lw
    wb = lw(workspace, data_only=False)
    values = lw(workspace, data_only=True)
    values[ROW_WISE]["D9"] = "15 900"
    markers = read_markers(wb[ROW_WISE], last_non_empty_row(values[ROW_WISE]))
    block = extract_sheet(values[ROW_WISE], wb[ROW_WISE], nomenclature, markers)[0]

    assert block.records[0].values["Premium"] == 15900.0
    notable = [c for c in block.coercions if not c.routine]
    assert [(c.field, c.source_ref) for c in notable] == [("Premium", "D9")]
    assert block.totals()["Premium"] == TOTALS["Premium"]


# ──────────────────────────────────────────── §9.2 S1, S2, S10 step-2 rules

def _relabel(block, labels):
    """Rewrite the Year of each record, keeping everything else intact."""
    for record, label in zip(block.records, labels):
        record.values["Year"] = label
    return block


def test_S1_natural_sort_keeps_same_year_variants_adjacent(books, nomenclature):
    b = _block(books, nomenclature, ROW_WISE)
    _relabel(b, ["2026 9 months", "2024", "2026", "2025", "2026 6 months", "2023"])
    result = apply_step2(b, step2_for(b.dataset.key))
    assert [r.values["Year"] for r in result.records] == [
        "2023", "2024", "2025", "2026", "2026 6 months", "2026 9 months",
    ]


def test_S1_natural_sort_orders_numerically_not_lexically(books, nomenclature):
    b = _block(books, nomenclature, ROW_WISE)
    _relabel(b, ["2021", "999", "2020", "10000", "2022", "37"])
    result = apply_step2(b, step2_for(b.dataset.key))
    assert [r.values["Year"] for r in result.records] == [
        "37", "999", "2020", "2021", "2022", "10000",
    ]


def test_S1_labels_are_never_rewritten(books, nomenclature):
    """The label is the identity — it must survive verbatim."""
    b = _block(books, nomenclature, ROW_WISE)
    _relabel(b, ["2026 9 months", "2024", "2025", "2026", "2023", "2022"])
    result = apply_step2(b, step2_for(b.dataset.key))
    assert "2026 9 months" in [r.values["Year"] for r in result.records]
    assert len(result.records) == 6, "records must not be merged by numeric year"


def test_S2_column_order_comes_from_sheet_00(books, nomenclature):
    b = _block(books, nomenclature, ROW_WISE)
    result = apply_step2(b, step2_for(b.dataset.key))
    assert result.declared_columns == nomenclature.datasets[ROW_WISE].headers
    assert result.columns == ("Year", "Premium", "Incurred Losses", "Loss Ratio %")


def test_S2_step2_spec_declares_no_column_order():
    """00 is the single source of truth for order — spec §9.2 S2."""
    assert not hasattr(step2_for("01 History"), "column_order")


def test_S10_overlapping_periods_raise_a_hypothesis(books, nomenclature):
    from datatransform.extract import _check_period_overlap, _number_hypotheses

    b = _block(books, nomenclature, ROW_WISE)
    _relabel(b, ["2024", "2025", "2026", "2026 9 months", "2027", "2028"])
    b.hypotheses.clear()
    _check_period_overlap(b)
    _number_hypotheses(b)

    assert len(b.hypotheses) == 1
    h = b.hypotheses[0]
    assert h.id == "H-01"
    assert h.attribute == "Period overlap 2026"
    assert h.value == "2026 · 2026 9 months"
    assert h.source == "tool"
    assert "extraction rather than being a portfolio total" in h.note


def test_S10_silent_when_every_period_is_distinct(books, nomenclature):
    from datatransform.extract import _check_period_overlap

    b = _block(books, nomenclature, ROW_WISE)
    _relabel(b, DISTINCT)
    b.hypotheses.clear()
    _check_period_overlap(b)
    assert b.hypotheses == []


def test_S10_flags_but_does_not_change_the_data(books, nomenclature):
    from datatransform.extract import _check_period_overlap

    b = _block(books, nomenclature, ROW_WISE)
    _relabel(b, ["2024", "2025", "2026", "2026 9 months", "2027", "2028"])
    before = [dict(r.values) for r in b.records]
    totals = b.totals()
    _check_period_overlap(b)
    assert [dict(r.values) for r in b.records] == before
    assert b.totals() == totals


def test_S10_reaches_the_written_block(workspace, tmp_path):
    """The reference workbook carries both a full year and a partial period."""
    out = tmp_path / "out.xlsx"
    report = run(workspace, out, tmp_path / "logs")
    assert report.ok

    ws = load_workbook(out, data_only=True)[ROW_WISE]
    text = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
    assert any("Period overlap 2025" in t for t in text)
    assert "2025 9 months" in text, "the label must reach the output verbatim"


# ════════════════════════════════════════ sheet 02 · globals · types · rules

EPI = "02. EPI Projections"
EPI_EXPECTED = [("2025 est", 18500.0), ("2025 9 months", 14400.0),
                ("2025 re-est", 19200.0), ("2026", 20900.0)]


def test_global_block_declares_N(nomenclature):
    assert nomenclature.actual_year == 2025
    assert nomenclature.globals["Cedent"] == "Example Insurance SA"


def test_types_come_from_sheet_00(nomenclature):
    assert nomenclature.type_of("Year") is FieldType.TEXT
    assert nomenclature.type_of("EPI") is FieldType.NUMBER
    assert nomenclature.type_of("Claim Reference") is FieldType.TEXT


def test_undeclared_type_is_fatal(nomenclature):
    with pytest.raises(ExtractionError, match="no datatype"):
        nomenclature.type_of("Something Undeclared")


def test_dataset_register_stops_at_the_next_block(nomenclature):
    """⟦RULES⟧ rows must not be read as datasets."""
    assert set(nomenclature.datasets) == {
        "01. History", "02. EPI Projections", "03. Large Losses",
        "04. Cat Losses", "05. Risk Profiles",
        "01. History_Transposed", "02. EPI Projections_Transposed",
    }


def test_rules_are_read_from_sheet_00(nomenclature):
    ids = [r.id for r in nomenclature.rules]
    assert ids == ["R-05", "R-06", "R-07", "R-01", "R-09", "R-02", "R-10"]
    assert [r.scope for r in nomenclature.rules] == [
        "all", "all", "all", "per risk", "per risk", "cat", "cat",
    ]
    r5 = nomenclature.rules[0]
    assert Rule.parse_ref(r5.left) == ("01 History", "Premium", "{N} 9 months")
    assert r5.relation == "=" and r5.tolerance == 1.0 and r5.severity == "error"


def test_epi_extraction(books, nomenclature):
    b = _block(books, nomenclature, EPI)
    assert b.address_map == {"Year": "B", "EPI": "C"}
    assert b.info_ref == "K"                       # a different selector column from 01
    assert (b.candidates, len(b.records), b.excluded) == (5, 4, 1)
    assert [(r.values["Year"], r.values["EPI"]) for r in b.records] == EPI_EXPECTED
    assert b.unextracted == ["D", "E"]             # Share % and Comment, not declared


def test_epi_overlap_flags_three_views_of_one_year(books, nomenclature):
    b = _block(books, nomenclature, EPI)
    overlap = [h for h in b.hypotheses if h.attribute == "Period overlap 2025"]
    assert len(overlap) == 1
    assert overlap[0].value == "2025 est · 2025 9 months · 2025 re-est"


# ───────────────────────────────────────────────── record matching

@pytest.mark.parametrize("pattern,expected", [
    ("{N} est", "2025 est"),
    ("{N} 9 months", "2025 9 months"),
    ("{N} re-est", "2025 re-est"),
    ("{N+1}", "2026"),
])
def test_records_match_on_leading_number_then_suffix(books, nomenclature, pattern, expected):
    b = _block(books, nomenclature, EPI)
    record = match_record(b, pattern, nomenclature.actual_year)
    assert record is not None and record.values["Year"] == expected


def test_bare_pattern_matches_a_lone_record_whatever_its_suffix(books, nomenclature):
    """{N+1} resolves whether the sheet writes 2026 or 2026 est."""
    b = _block(books, nomenclature, EPI)
    b.records[-1].values["Year"] = "2026 est"
    assert match_record(b, "{N+1}", 2025).values["Year"] == "2026 est"


def test_ambiguous_bare_pattern_matches_nothing(books, nomenclature):
    """Where several records share the year, the suffix must be exact."""
    b = _block(books, nomenclature, EPI)
    assert match_record(b, "{N}", 2025) is None      # est, 9 months and re-est all exist


def test_split_label():
    assert split_label("2025 9 months") == (2025, "9 months")
    assert split_label("2026") == (2026, "")
    assert split_label("Total") == (None, "Total")


# ───────────────────────────────────────────────── derived figures

def test_derived_figures_are_block_level(books, nomenclature):
    b = _block(books, nomenclature, EPI)
    result = apply_step2(b, step2_for(b.dataset.key), nomenclature)
    figures = {f.name: f.value for f in result.figures}
    assert figures["Estimation error"] == pytest.approx(19200 / 18500 - 1)
    assert figures["Implied growth"] == pytest.approx(20900 / 19200 - 1)


def test_derived_figures_report_rather_than_guess(books, nomenclature):
    b = _block(books, nomenclature, EPI)
    b.records = [r for r in b.records if r.values["Year"] != "2025 est"]
    result = apply_step2(b, step2_for(b.dataset.key), nomenclature)
    estimation = next(f for f in result.figures if f.name == "Estimation error")
    assert estimation.value is None
    assert "no record matching" in estimation.detail


def test_figures_need_N(books, nomenclature):
    b = _block(books, nomenclature, EPI)
    result = apply_step2(b, step2_for(b.dataset.key), nomenclature=None)
    assert all(f.value is None for f in result.figures)
    assert all("Actual year" in f.detail for f in result.figures)


# ───────────────────────────────────────────────── crosschecks

def _blocks(books, nomenclature, sheets=(ROW_WISE, EPI, "03. Large Losses")):
    return {
        b.dataset.key: b
        for sheet in sheets
        for b in [_block(books, nomenclature, sheet)]
    }


def _find(blocks, key):
    return blocks[key]


def test_all_declared_rules_pass_on_the_reference_workbook(books, nomenclature):
    """Everything that applies passes; the cat rule is not applicable on a Fire pack."""
    results = run_rules(nomenclature, _blocks(books, nomenclature))
    assert {r.status for r in results} == {"passed", "not applicable"}
    # Both cat rules stay silent on a Fire pack: structure, not a gap.
    assert [r.rule.id for r in results if r.status == "not applicable"] == ["R-02", "R-10"]
    assert results[0].left_value == results[0].right_value == 14400.0
    assert results[1].left_value == results[1].right_value == 19200.0


def test_a_broken_tie_fails_the_rule(books, nomenclature):
    blocks = _blocks(books, nomenclature)
    epi = blocks["02 EPI"]
    next(r for r in epi.records if r.values["Year"] == "2025 9 months").values["EPI"] = 13000.0
    results = run_rules(nomenclature, blocks)
    assert results[0].status == "failed"
    assert "14,400 = 13,000" in results[0].detail


def test_differing_premium_basis_skips_rather_than_compares(books, nomenclature):
    """Comparing GWP against GNPI silently would be worse than not checking."""
    from datatransform.model import Attribute

    blocks = _blocks(books, nomenclature)
    blocks["02 EPI"].attributes["Premium basis"] = Attribute("Premium basis", "GWP", False, 4)
    results = run_rules(nomenclature, blocks)
    assert results[0].status == "skipped"
    assert "Premium basis differs (GNPI vs GWP)" in results[0].detail


def test_scale_is_normalised_before_comparing(books, nomenclature):
    """01 in thousands against 02 in units must still tie."""
    from datatransform.model import Attribute

    blocks = _blocks(books, nomenclature)
    epi = blocks["02 EPI"]
    epi.attributes["Scale"] = Attribute("Scale", "1", False, 3)
    for record in epi.records:
        record.values["EPI"] *= 1000
    results = run_rules(nomenclature, blocks)
    assert {r.status for r in results} <= {"passed", "not applicable"}


def test_a_missing_dataset_is_not_applicable_rather_than_a_failure(books, nomenclature):
    """A pack without that sheet is structure, not a problem — spec §10.1."""
    blocks = {"02 EPI": _block(books, nomenclature, EPI)}
    results = run_rules(nomenclature, blocks)
    by_id = {r.rule.id: r for r in results}
    assert by_id["R-05"].status == "not applicable"
    assert "not in this pack" in by_id["R-05"].detail
    assert by_id["R-07"].status == "passed"       # R-07 is within 02 alone
    assert hypotheses_from(results, blocks["02 EPI"]) == []


def test_failed_and_skipped_rules_become_questions(books, nomenclature):
    blocks = _blocks(books, nomenclature)
    next(r for r in blocks["02 EPI"].records
         if r.values["Year"] == "2025 9 months").values["EPI"] = 13000.0
    results = run_rules(nomenclature, blocks)
    hyps = hypotheses_from(results, blocks["02 EPI"])
    assert [h.attribute for h in hyps] == ["Crosscheck R-05/Fire"]
    assert hyps[0].confidence is Confidence.OPEN


def test_crosschecks_reach_both_sheets(workspace, tmp_path):
    out = tmp_path / "out.xlsx"
    report = run(workspace, out, tmp_path / "logs")
    assert report.ok and report.rules_ok

    wb = load_workbook(out, data_only=True)
    for sheet in (ROW_WISE, EPI):
        text = [c.value for row in wb[sheet].iter_rows() for c in row
                if isinstance(c.value, str)]
        assert any("Crosschecks against other sheets" in t for t in text), sheet
        assert any(t.startswith("R-05") for t in text), sheet


def test_derived_figures_are_written_to_the_sheet(workspace, tmp_path):
    out = tmp_path / "out.xlsx"
    run(workspace, out, tmp_path / "logs")
    ws = load_workbook(out, data_only=True)[EPI]
    labels = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
    assert "Estimation error" in labels
    assert "Implied growth" in labels

    values = [c.value for row in ws.iter_rows() for c in row
              if isinstance(c.value, float) and 0.03 < c.value < 0.09]
    assert any(abs(v - (19200 / 18500 - 1)) < 1e-9 for v in values)


# ════════════════════════════════════ ⟦PERIOD ORDER⟧ · transposed 02

EPI_T = "02. EPI Projections_Transposed"
LARGE = "03. Large Losses"


def test_period_order_is_read_from_sheet_00(nomenclature):
    assert nomenclature.period_order == {"est": 1, "9 months": 2, "re-est": 3}


def test_declared_period_order_beats_alphanumeric(books, nomenclature):
    """'9 months' sorts after 'est', though '9' < 'e' alphanumerically."""
    b = _block(books, nomenclature, EPI)
    result = apply_step2(b, step2_for(b.dataset.key), nomenclature)
    assert [r.values["Year"] for r in result.records] == [
        "2025 est", "2025 9 months", "2025 re-est", "2026",
    ]


def test_without_a_declared_order_it_falls_back_to_alphanumeric(books, nomenclature):
    from datatransform.nomenclature import Nomenclature

    b = _block(books, nomenclature, EPI)
    bare = Nomenclature(nomenclature.datasets, nomenclature.vocabulary,
                        nomenclature.globals, nomenclature.types, nomenclature.rules)
    result = apply_step2(b, step2_for(b.dataset.key), bare)
    assert [r.values["Year"] for r in result.records] == [
        "2025 9 months", "2025 est", "2025 re-est", "2026",
    ]


def test_a_bare_year_sorts_before_its_qualified_views(books, nomenclature):
    b = _block(books, nomenclature, ROW_WISE)
    result = apply_step2(b, step2_for(b.dataset.key), nomenclature)
    years = [r.values["Year"] for r in result.records]
    assert years[-2:] == ["2025", "2025 9 months"]


def test_an_undeclared_suffix_sorts_last_and_stays_deterministic(books, nomenclature):
    b = _block(books, nomenclature, EPI)
    b.records[-1].values["Year"] = "2025 half year"      # not in ⟦PERIOD ORDER⟧
    result = apply_step2(b, step2_for(b.dataset.key), nomenclature)
    assert [r.values["Year"] for r in result.records] == [
        "2025 est", "2025 9 months", "2025 re-est", "2025 half year",
    ]


def test_sort_note_records_the_period_order(books, nomenclature):
    b = _block(books, nomenclature, EPI)
    result = apply_step2(b, step2_for(b.dataset.key), nomenclature)
    assert "period order est → 9 months → re-est" in result.notes[0]


# ───────────────────────────────────────── transposed 02

def test_transposed_epi_resolution(books, nomenclature):
    b = _block(books, nomenclature, EPI_T)
    assert b.orientation is Orientation.TRANSPOSED
    assert b.header_ref == "C" and b.info_ref == "14"
    assert b.address_map == {"Year": "8", "EPI": "10"}
    assert (b.candidates, len(b.records), b.excluded) == (5, 4, 1)


def test_both_epi_orientations_are_equivalent(books, nomenclature):
    a = _block(books, nomenclature, EPI)
    b = _block(books, nomenclature, EPI_T)
    assert [r.values for r in a.records] == [r.values for r in b.records]
    assert a.totals() == b.totals() == {"EPI": 73000.0}
    assert a.provenance_label != b.provenance_label


def test_derived_figures_match_across_orientations(books, nomenclature):
    figures = {}
    for sheet in (EPI, EPI_T):
        b = _block(books, nomenclature, sheet)
        result = apply_step2(b, step2_for(b.dataset.key), nomenclature)
        figures[sheet] = {f.name: f.value for f in result.figures}
    assert figures[EPI] == figures[EPI_T]
    assert figures[EPI]["Estimation error"] == pytest.approx(19200 / 18500 - 1)


def test_record_matching_works_transposed(books, nomenclature):
    b = _block(books, nomenclature, EPI_T)
    assert match_record(b, "{N} re-est", 2025).values["EPI"] == 19200.0
    assert match_record(b, "{N+1}", 2025).values["EPI"] == 20900.0


def test_all_data_sheets_process(workspace, tmp_path):
    out = tmp_path / "out.xlsx"
    report = run(workspace, out, tmp_path / "logs")
    assert report.ok and report.rules_ok
    processed = {o.sheet: o.detail for o in report.outcomes if o.status == "processed"}
    assert set(processed) == {ROW_WISE, TRANSPOSED, EPI, EPI_T, LARGE}
    assert processed[EPI] == processed[EPI_T] == "1 block(s), 4 record(s)"
    assert processed[ROW_WISE] == processed[TRANSPOSED] == "1 block(s), 6 record(s)"


# ══════════════════════════════════ 03. Large Losses · dates · aggregation

LOSS_TOTAL = 12760.0
ANNUAL = [("2021", 2770.0), ("2022", 0.0), ("2023", 5330.0),
          ("2024", 2600.0), ("2025", 2060.0)]


def test_date_type_is_declared_in_sheet_00(nomenclature):
    assert nomenclature.type_of("Date of Loss") is FieldType.DATE
    assert FieldType.DATE.number_format == "yyyy-mm-dd"


@pytest.mark.parametrize("raw,expected", [
    ("2023-06-15", date(2023, 6, 15)),
    ("2023/06/15", date(2023, 6, 15)),
    ("15.06.2023", date(2023, 6, 15)),      # 15 cannot be a month
    ("15/06/2023", date(2023, 6, 15)),
    (datetime(2024, 5, 19, 0, 0), date(2024, 5, 19)),
    (date(2024, 5, 19), date(2024, 5, 19)),
    (None, None),
    ("", None),
])
def test_date_coercion(raw, expected):
    assert to_date(raw, "Date of Loss", "F11", []) == expected


def test_ambiguous_date_is_fatal():
    """01/02/2025 is 1 February or 2 January — never guess."""
    with pytest.raises(ExtractionError, match="ambiguous"):
        to_date("01/02/2025", "Date of Loss", "F11", [])


def test_a_declared_date_format_resolves_the_ambiguity():
    assert to_date("01/02/2025", "Date of Loss", "F11", [], "DD.MM.YYYY") == date(2025, 2, 1)
    assert to_date("01/02/2025", "Date of Loss", "F11", [], "MM/DD/YYYY") == date(2025, 1, 2)


def test_impossible_date_is_fatal():
    with pytest.raises(ExtractionError, match="not a real date"):
        to_date("2023-02-30", "Date of Loss", "F11", [])


def test_large_losses_extraction(books, nomenclature):
    b = _block(books, nomenclature, LARGE)
    assert b.address_map == {"Year": "C", "Name of Loss": "D",
                             "Loss amount": "E", "Date of Loss": "F"}
    assert (b.candidates, len(b.records), b.excluded) == (9, 8, 1)
    assert b.numeric_fields == ("Loss amount",)          # the date is never summed
    assert b.totals()["Loss amount"] == LOSS_TOTAL
    first = b.records[0].values
    assert first["Year"] == "2021"
    assert first["Name of Loss"] == "Warehouse fire, Lyon"
    assert first["Date of Loss"] == date(2021, 3, 14)


def test_many_claims_in_one_year_is_not_a_period_overlap(books, nomenclature):
    """Eight losses across four years are claims, not overlapping periods."""
    b = _block(books, nomenclature, LARGE)
    assert not [h for h in b.hypotheses if h.attribute.startswith("Period overlap")]

    history = _block(books, nomenclature, ROW_WISE)      # but a real overlap still fires
    assert [h.attribute for h in history.hypotheses
            if h.attribute.startswith("Period overlap")] == ["Period overlap 2025"]


# ───────────────────────────────────────────────── aggregation

def _step2_large(books, nomenclature):
    b = _block(books, nomenclature, LARGE)
    return b, apply_step2(b, step2_for(b.dataset.key), nomenclature,
                          _blocks(books, nomenclature))


def test_aggregate_groups_by_year_and_sums(books, nomenclature):
    _, result = _step2_large(books, nomenclature)
    table = result.aggregates[0]
    assert table.title == "Annual sum of large losses"
    assert [(k, v["Loss amount"]) for k, v in table.rows] == ANNUAL


def test_a_year_without_losses_shows_zero(books, nomenclature):
    """Absent reads as 'no data'; 0 reads as 'nothing happened'."""
    _, result = _step2_large(books, nomenclature)
    assert result.aggregates[0].zero_filled == ["2022"]
    assert dict(result.aggregates[0].rows)["2022"]["Loss amount"] == 0.0


def test_aggregate_total_ties_to_the_detail(books, nomenclature):
    _, result = _step2_large(books, nomenclature)
    assert result.aggregates[0].totals()["Loss amount"] == result.totals()["Loss amount"]
    assert result.aggregates[0].totals()["Loss amount"] == LOSS_TOTAL


def test_grouping_that_loses_records_is_fatal(books, nomenclature, monkeypatch):
    from datatransform import transform as tmod

    b = _block(books, nomenclature, LARGE)
    real = tmod._aggregate
    monkeypatch.setattr(tmod, "_aggregate",
                        lambda blk, spec, blocks: _drop_row(real(blk, spec, blocks)))
    with pytest.raises(ExtractionError, match="grouping lost records"):
        apply_step2(b, step2_for(b.dataset.key), nomenclature, _blocks(books, nomenclature))


def _drop_row(table):
    table.rows = table.rows[:-1]
    return table


def test_detail_sorts_by_year_then_date(books, nomenclature):
    _, result = _step2_large(books, nomenclature)
    dates = [r.values["Date of Loss"] for r in result.records]
    assert dates == sorted(dates)


# ───────────────────────────────────────── reference grammar & rules

@pytest.mark.parametrize("ref,key,name,record,agg", [
    ("01 History.Premium@{N}", "01 History", "Premium", "{N}", None),
    ("01 History.Year basis", "01 History", "Year basis", None, None),
    ("SUM(03 Large.Loss amount@{Y})", "03 Large", "Loss amount", "{Y}", "SUM"),
])
def test_reference_forms(ref, key, name, record, agg):
    r = parse_reference(ref)
    assert (r.dataset_key, r.name, r.record, r.aggregate) == (key, name, record, agg)
    assert r.is_attribute == (record is None)


def test_a_malformed_reference_is_fatal():
    with pytest.raises(ExtractionError, match="is not of the form"):
        parse_reference("nonsense")


def test_attribute_rule_compares_year_basis(books, nomenclature):
    results = run_rules(nomenclature, _blocks(books, nomenclature))
    r9 = next(r for r in results if r.rule.id == "R-09")
    assert r9.status == "passed" and r9.detail == "UW = UW"


def test_a_differing_year_basis_fails_the_attribute_rule(books, nomenclature):
    from datatransform.model import Attribute

    blocks = _blocks(books, nomenclature)
    blocks["03 Large"].attributes["Year basis"] = Attribute("Year basis", "Occurrence", True, 5)
    results = run_rules(nomenclature, blocks)
    r9 = next(r for r in results if r.rule.id == "R-09")
    assert r9.status == "failed" and r9.detail == "UW = Occurrence"


def test_year_wildcard_expands_one_rule_per_year(books, nomenclature):
    results = run_rules(nomenclature, _blocks(books, nomenclature))
    r1 = [r for r in results if r.rule.id == "R-01"]
    assert [r.year for r in r1] == [2021, 2023, 2024, 2025]
    assert [r.label for r in r1] == ["R-01/Fire/2021", "R-01/Fire/2023",
                                     "R-01/Fire/2024", "R-01/Fire/2025"]
    assert all(r.status == "passed" for r in r1)


def test_sum_aggregates_the_years_losses(books, nomenclature):
    results = run_rules(nomenclature, _blocks(books, nomenclature))
    r2023 = next(r for r in results if r.rule.id == "R-01" and r.year == 2023)
    assert r2023.left_value == 5330.0            # three losses summed
    assert r2023.right_value == 12660.0


def test_large_losses_exceeding_incurred_fails(books, nomenclature):
    blocks = _blocks(books, nomenclature)
    for record in blocks["03 Large"].records:
        record.values["Loss amount"] *= 10
    results = run_rules(nomenclature, blocks)
    failed = [r for r in results if r.status == "failed"]
    assert [r.label for r in failed] == ["R-01/Fire/2021", "R-01/Fire/2023",
                                         "R-01/Fire/2024", "R-01/Fire/2025"]
    assert hypotheses_from(results, blocks["03 Large"])[0].confidence is Confidence.OPEN


def test_differing_share_basis_skips_the_comparison(books, nomenclature):
    from datatransform.model import Attribute

    blocks = _blocks(books, nomenclature)
    blocks["03 Large"].attributes["Share basis"] = Attribute("Share basis", "ceded only", False, 4)
    results = run_rules(nomenclature, blocks)
    r1 = next(r for r in results if r.rule.id == "R-01")
    assert r1.status == "skipped"
    assert "Share basis differs (ceded only vs 100%)" in r1.detail   # 03 is the left side


# ───────────────────────────────────── occurrence-year consistency

def test_occurrence_basis_flags_a_year_that_disagrees_with_the_date(books, nomenclature):
    from datatransform.extract import _check_occurrence_year
    from datatransform.model import Attribute

    b = _block(books, nomenclature, LARGE)
    b.attributes["Year basis"] = Attribute("Year basis", "Occurrence", True, 5)
    b.records[0].values["Year"] = "2019"          # date is 2021-03-14
    b.hypotheses.clear()
    _check_occurrence_year(b)
    assert len(b.hypotheses) == 1
    assert b.hypotheses[0].attribute == "Occurrence year mismatch"
    assert b.hypotheses[0].confidence is Confidence.OPEN


def test_underwriting_basis_stays_silent(books, nomenclature):
    """Under a UW basis a loss may legitimately fall in a later year."""
    from datatransform.extract import _check_occurrence_year

    b = _block(books, nomenclature, LARGE)
    b.records[0].values["Year"] = "2019"
    b.hypotheses.clear()
    _check_occurrence_year(b)                     # declared basis is UW
    assert b.hypotheses == []


def test_aggregate_reaches_the_written_sheet(workspace, tmp_path):
    out = tmp_path / "out.xlsx"
    report = run(workspace, out, tmp_path / "logs")
    assert report.ok and report.rules_ok

    ws = load_workbook(out, data_only=True)[LARGE]
    text = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
    assert "Annual sum of large losses" in text
    assert any("shown as 0: 2022" in t for t in text)
    assert any(t.startswith("R-01/Fire/2021") for t in text)


def test_dates_are_written_as_dates(workspace, tmp_path):
    out = tmp_path / "out.xlsx"
    run(workspace, out, tmp_path / "logs")
    ws = load_workbook(out, data_only=True)[LARGE]
    written = [c for row in ws.iter_rows() for c in row
               if isinstance(c.value, datetime) and c.number_format == "yyyy-mm-dd"]
    assert len(written) == 24            # eight losses × source, step 1 and step 2
