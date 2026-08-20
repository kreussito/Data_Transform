#!/usr/bin/env python3
"""Reusable sheet builders for the reference workbooks.

Every generated workbook is assembled from the same pieces, so the four treaty shapes
differ only in their ⟦SECTIONS⟧ block and which sheets they carry.
"""

from __future__ import annotations

from openpyxl.styles import Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as col_letter

FONT = "Arial"
title_f = Font(name=FONT, size=14, bold=True)
sub_f = Font(name=FONT, size=9, italic=True, color="666666")
head_f = Font(name=FONT, size=10, bold=True)
body_f = Font(name=FONT, size=10)
inert_f = Font(name=FONT, size=10, italic=True, color="999999")
marker_f = Font(name=FONT, size=10, bold=True, color="7F6000")
attr_f = Font(name=FONT, size=10, color="7F6000")
prov_f = Font(name=FONT, size=10, italic=True, color="999999")
mono_f = Font(name="Consolas", size=9)

yellow = PatternFill("solid", fgColor="FFF2CC")   # human-authored markers
green = PatternFill("solid", fgColor="E2EFDA")    # selector formulas
blue = PatternFill("solid", fgColor="DDEBF7")     # extraction row / column
grey = PatternFill("solid", fgColor="F2F2F2")     # block headers

thin = Side(style="thin", color="BFBFBF")
box = Border(left=thin, right=thin, top=thin, bottom=thin)

STRUCTURAL = ("Header_", "Info_", "Transpose_")


def put(ws, ref, value, font=body_f, fill=None, fmt=None):
    cell = ws[ref]
    cell.value = value
    cell.font = font
    if fill:
        cell.fill = fill
    if fmt:
        cell.number_format = fmt
    return cell


def block_header(ws, row, anchor, columns):
    """An ⟦ANCHOR⟧ header followed by a column-label row."""
    put(ws, f"B{row}", anchor, head_f, grey)
    for i, label in enumerate(columns):
        put(ws, f"{col_letter(2 + i)}{row + 1}", label, head_f)
    return row + 2


def markers(ws, entries: dict):
    """Write column A. Structural tokens are bold; attributes are plain."""
    for row, text in entries.items():
        cell = ws.cell(row=row, column=1, value=text)
        cell.fill = yellow
        cell.font = marker_f if text.startswith(STRUCTURAL) else attr_f
        cell.border = box


def record_block(ws, *, header_row, selector_col, source_labels, declared, rows,
                 formats, note_col="H", total_cols=None):
    """One row-wise block: an inert source header, an extraction row, then records.

    ``declared`` maps a column letter to the label declared in sheet 00; every other
    column of the source is deliberately left out of the extraction row.
    """
    for col, text in source_labels.items():
        put(ws, f"{col}{header_row - 1}", text, inert_f)
    put(ws, f"{note_col}{header_row - 1}", "← source header, never read", sub_f)

    for col, label in declared.items():
        put(ws, f"{col}{header_row}", label, head_f, blue).border = box
    put(ws, f"{note_col}{header_row}",
        "← extraction row: blanks mean 'not extracted'", sub_f)

    first = header_row + 1
    for i, values in enumerate(rows):
        row = first + i
        for col, value in values.items():
            put(ws, f"{col}{row}", value, body_f, fmt=formats.get(col))
        cell = ws[f"{selector_col}{row}"]
        cell.value = "=ROW()"
        cell.font, cell.fill, cell.border = body_f, green, box
    last = first + len(rows) - 1

    total_row = last + 1
    # ``total_cols`` names the columns the cedent would actually total. A band bound is
    # a number but not a quantity, so a profile's own total row leaves the bounds alone.
    numeric = (total_cols if total_cols is not None
               else [c for c, f in formats.items() if f == "#,##0"])
    key_col = next(iter(declared))
    put(ws, f"{key_col}{total_row}", "Total", head_f)
    for col in numeric:
        cell = ws[f"{col}{total_row}"]
        cell.value = f"=SUM({col}{first}:{col}{last})"
        cell.font, cell.number_format = head_f, "#,##0"
    put(ws, f"{note_col}{total_row}",
        "← selector empty: excluded from extraction, reused as control", sub_f)
    return total_row


def widths(ws, spec: dict):
    for col, width in spec.items():
        ws.column_dimensions[col].width = width


# ── the dataset inventory ─────────────────────────────────────────────────────
# The register above it says what *this pack* carries. This block says what the
# standard defines, so a reader can tell a sheet that is absent from a sheet that
# was never specified — spec §2.
#
#   role · sheet · what it holds · state
INVENTORY = [
    ("01", "History", "Premium and loss history per section", "implemented"),
    ("02", "EPI Projections", "Estimated premium income, N and N+1", "implemented"),
    ("03", "Large Losses", "Individual large claims, per-risk sections", "implemented"),
    ("04", "Cat Losses", "Catastrophe events, cat sections", "implemented"),
    ("05", "Risk Profiles", "Banded exposure — the per-risk rating basis", "implemented"),
    ("06", "EQ Aggs", "Earthquake sums insured per cat zone", "implemented"),
    ("07", "Wind Aggs", "Windstorm sums insured per cat zone", "implemented"),
    ("08", "Splits", "Occupancy x cover (Fire) or x Projects/Renewables (Eng.)",
     "axes settled, measures open"),
    ("09", "Rate Development", "Rate change history — optional", "outstanding"),
    ("10", "Triangles", "Loss development triangles", "outstanding"),
    ("11", "Exchange rates", "FX rates for conversion between currencies", "outstanding"),
    ("20", "Summary", "Collected step-2 blocks — written, never read", "outstanding"),
]


def inventory_block(ws, row, note=None):
    """Write ⟦INVENTORY⟧ — every dataset the standard defines, not just this pack's."""
    r = block_header(ws, row, "⟦INVENTORY⟧", ["Role", "Sheet", "Holds", "State"])
    for role, sheet, holds, state in INVENTORY:
        put(ws, f"B{r}", role, body_f)
        put(ws, f"C{r}", sheet, body_f)
        put(ws, f"D{r}", holds, body_f)
        put(ws, f"E{r}", state, body_f, grey if state != "implemented" else None)
        r += 1
    put(ws, f"B{r + 1}", note or (
        "Every dataset the tool knows, whether or not this pack carries it. The register "
        "above lists what is here; ⟦SECTIONS⟧ says which roles each section expects, so a "
        "missing sheet reads as structure rather than as a gap."), sub_f)
    return r + 2
