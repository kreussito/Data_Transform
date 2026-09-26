"""Every dataset on its side — specification_v1.md §6.2.

`01` and `02` have carried transposed twins since the beginning, and until now that was
the whole of the evidence: whether `03`–`09` survived the flip was an assumption. This
file removes it, by flipping a finished pack mechanically and demanding the same answer.

Nothing in the transposed pack is retyped — ``tools/transpose_pack.py`` moves the cells
across — so any difference in the output belongs to the tool and not to the fixture.
"""

from __future__ import annotations

import shutil
from pathlib import Path

import pytest
from openpyxl import load_workbook

from datatransform.extract import extract_sheet, last_non_empty_row
from datatransform.markers import read_markers
from datatransform.model import Orientation
from datatransform.nomenclature import read_nomenclature
from datatransform.runner import run
from datatransform.specs import step2_for
from datatransform.transform import apply_step2

ROOT = Path(__file__).resolve().parents[1]
ROW_WISE = ROOT / "Intake_FireCatFull_v1.xlsx"
TRANSPOSED = ROOT / "Intake_FireCatFull_Transposed_v1.xlsx"


def _read(path):
    values = load_workbook(path, data_only=True)
    formulas = load_workbook(path, data_only=False)
    nomenclature = read_nomenclature(formulas)
    blocks = []
    for title in formulas.sheetnames:
        if title.startswith("00."):
            continue
        markers = read_markers(formulas[title], last_non_empty_row(values[title]))
        blocks.extend(extract_sheet(values[title], formulas[title], nomenclature, markers))
    return nomenclature, blocks


def _keyed(blocks):
    return {(b.dataset.key, b.section, b.index): b for b in blocks}


@pytest.fixture(scope="module")
def packs():
    row_nom, row_blocks = _read(ROW_WISE)
    col_nom, col_blocks = _read(TRANSPOSED)
    return (row_nom, row_blocks, _keyed(row_blocks)), (col_nom, col_blocks,
                                                       _keyed(col_blocks))


# ─────────────────────────────────────────── every dataset, not just 01 and 02

def test_every_dataset_is_present_on_its_side(packs):
    (_, _, by_row), (_, _, by_col) = packs
    assert set(by_row) == set(by_col)
    assert {k[0].split()[0] for k in by_col} == {
        "01", "02", "03", "04", "05", "06", "07", "08", "09"}


def test_they_really_are_transposed(packs):
    (_, _, by_row), (_, _, by_col) = packs
    assert all(b.orientation is Orientation.ROW_WISE for b in by_row.values())
    assert all(b.orientation is Orientation.TRANSPOSED for b in by_col.values())
    assert all(b.provenance_label == "Source column" for b in by_col.values())


def test_step_1_is_identical_record_for_record(packs):
    (_, _, by_row), (_, _, by_col) = packs
    for key, block in sorted(by_row.items()):
        twin = by_col[key]
        assert block.fields == twin.fields, key
        assert len(block.records) == len(twin.records), key
        for a, b in zip(block.records, twin.records):
            assert a.values == b.values, f"{key} record {a.source_ref}/{b.source_ref}"


def test_step_2_reaches_the_same_totals(packs):
    (row_nom, row_blocks, by_row), (col_nom, col_blocks, by_col) = packs
    for key, block in sorted(by_row.items()):
        spec = step2_for(block.dataset.key)
        if not spec:
            continue
        mine = apply_step2(block, spec, row_nom, row_blocks)
        theirs = apply_step2(by_col[key], spec, col_nom, col_blocks)
        assert mine.totals() == theirs.totals(), key
        assert [r.values for r in mine.records] == [r.values for r in theirs.records], key


# ────────────────────────── several blocks sharing one label column

def test_a_stacked_transposed_sheet_finds_each_block_its_own_labels():
    """Row-wise, each block has its own extraction *row*, so a scan of it can never pick
    up a neighbour's labels. Transposed, three versions share one label *column* and
    ``Zone`` appears three times in it. The band comes from the Info markers, which are
    already in the sheet — nothing new is declared."""
    _, blocks = _read(TRANSPOSED)
    aggs = sorted((b for b in blocks if b.dataset.role == "06"), key=lambda b: b.index)

    assert len(aggs) == 3
    assert {b.header_ref for b in aggs} == {"C"}          # one shared label column
    assert len({b.info_ref for b in aggs}) == 3           # three selector rows
    for block in aggs:
        assert "Zone" in block.address_map
    assert len({b.address_map["Zone"] for b in aggs}) == 3


def test_the_bands_do_not_overlap():
    from datatransform.extract import _label_band

    values = load_workbook(TRANSPOSED, data_only=True)["06. EQ Aggs_T"]
    formulas = load_workbook(TRANSPOSED, data_only=False)["06. EQ Aggs_T"]
    markers = read_markers(formulas, last_non_empty_row(values))

    bands = [_label_band(markers, i, values.max_row) for i in (1, 2, 3)]
    for (first, last), (next_first, _) in zip(bands, bands[1:]):
        assert first <= last < next_first


def test_09_transposed_keeps_its_two_scopes():
    _, blocks = _read(TRANSPOSED)
    rates = sorted((b for b in blocks if b.dataset.role == "09"), key=lambda b: b.index)
    assert [b.attributes["Scope"].value for b in rates] == [
        "Fire", "Earthquake + Windstorm"]


# ───────────────────────────────────────────────────────── end to end

def test_the_transposed_pack_runs_clean(tmp_path):
    source = tmp_path / TRANSPOSED.name
    shutil.copy2(TRANSPOSED, source)
    report = run(source, tmp_path / "out.xlsx", tmp_path / "logs")

    assert report.ok, [o.detail for o in report.outcomes if o.status == "error"]
    assert report.rules_ok
    assert len({o.sheet for o in report.outcomes if o.status == "processed"}) == 16


def test_the_group_checks_reach_the_same_answers(tmp_path):
    """§10.3, §10.4 and §10.5 read blocks, not layouts — so the orientation must not
    reach them at all."""
    outcomes = {}
    for path in (ROW_WISE, TRANSPOSED):
        source = tmp_path / path.name
        shutil.copy2(path, source)
        report = run(source, tmp_path / f"{path.stem}_out.xlsx", tmp_path / "logs")
        outcomes[path.stem] = report

    row, col = outcomes[ROW_WISE.stem], outcomes[TRANSPOSED.stem]
    assert [(t.section, t.exposure_growth, t.implied_rate_change) for t in row.growth] == \
        [(t.section, t.exposure_growth, t.implied_rate_change) for t in col.growth]
    assert [(c.scope, c.claimed, c.implied, c.status) for c in row.rate_claims] == \
        [(c.scope, c.claimed, c.implied, c.status) for c in col.rate_claims]
