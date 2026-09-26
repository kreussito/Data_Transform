"""Declared losses against total incurred — specification_v1.md §10.3."""

from __future__ import annotations

import shutil
from pathlib import Path

import pytest
from openpyxl import load_workbook

from datatransform.extract import extract_sheet, last_non_empty_row
from datatransform.lossshare import (
    DEFAULT_THRESHOLD,
    EXCEEDS,
    OK,
    WARNING,
    loss_share_tables,
    threshold_of,
)
from datatransform.markers import read_markers
from datatransform.model import Attribute
from datatransform.nomenclature import read_nomenclature
from datatransform.runner import run

ROOT = Path(__file__).resolve().parents[1]

FIRE = ROOT / "Intake_v1.xlsx"
ENGINEERING = ROOT / "Intake_Engineering_v1.xlsx"
FIRE_CAT = ROOT / "Intake_FireCat_v1.xlsx"
FIRE_EQ_WIND = ROOT / "Intake_FireEQWind_v1.xlsx"


def _read(path):
    values = load_workbook(path, data_only=True)
    formulas = load_workbook(path, data_only=False)
    nomenclature = read_nomenclature(formulas)
    blocks = []
    for title in formulas.sheetnames:
        if title.startswith("00."):
            continue
        markers = read_markers(formulas[title], last_non_empty_row(values[title]))
        blocks.extend(extract_sheet(values[title], formulas[title], nomenclature, markers))
    return nomenclature, blocks


def _tables(path):
    nomenclature, blocks = _read(path)
    return {t.kind: t for t in loss_share_tables(nomenclature, blocks)}


def _row(table, year):
    return next(r for r in table.rows if r.year == year)


# ───────────────────────────────────────── the three cases the user named

def test_engineering_sums_large_and_cat_losses():
    """Both exist, so both are added — spec §10.3."""
    table = _tables(ENGINEERING)["per risk"]
    assert table.roles == ("03", "04")
    assert table.sections == ("Engineering",)

    row = _row(table, 2024)
    assert row.by_role == {"03": 1975.0, "04": 2600.0}
    assert row.declared == 4575.0
    assert row.incurred == 7240.0


def test_where_only_one_loss_dataset_exists_that_one_is_taken():
    fire = _tables(FIRE)["per risk"]
    assert fire.roles == ("03",)                       # Fire has no cat sheet
    assert _row(fire, 2023).declared == 5330.0

    cat = _tables(FIRE_EQ_WIND)["cat"]
    assert cat.roles == ("04",)                        # cat sections have no large losses


def test_the_cat_sections_are_summed_together():
    """Earthquake and Windstorm are one cat portfolio, checked as one."""
    table = _tables(FIRE_EQ_WIND)["cat"]
    assert table.sections == ("Earthquake", "Windstorm")
    # 2021: EQ 950 + Wind 1,400 against EQ 1,200 + Wind 2,300
    row = _row(table, 2021)
    assert row.declared == 2350.0
    assert row.incurred == 3500.0


def test_per_risk_and_cat_are_checked_separately():
    tables = _tables(FIRE_EQ_WIND)
    assert set(tables) == {"per risk", "cat"}
    assert tables["per risk"].sections == ("Fire",)
    assert tables["cat"].sections == ("Earthquake", "Windstorm")
    # Fire's incurred is not in the cat group's basis and vice versa.
    assert _row(tables["per risk"], 2021).incurred == 14930.0
    assert _row(tables["cat"], 2021).incurred == 3500.0


def test_the_two_cat_shapes_produce_the_same_check():
    """One sheet or three — the check is about sections, not sheets."""
    combined, separate = _tables(FIRE_CAT)["cat"], _tables(FIRE_EQ_WIND)["cat"]
    assert [(r.year, r.declared, r.incurred) for r in combined.rows] == \
           [(r.year, r.declared, r.incurred) for r in separate.rows]


# ────────────────────────────────────────────────── the two questions

def test_a_year_within_both_limits_says_so():
    row = _row(_tables(FIRE)["per risk"], 2021)
    assert row.share == pytest.approx(2770 / 14930)     # 18.6%
    assert row.status(DEFAULT_THRESHOLD) == OK


def test_a_share_above_the_threshold_warns():
    table = _tables(FIRE)["per risk"]
    assert _row(table, 2023).status(table.threshold) == WARNING   # 42.1%
    assert _row(table, 2025).status(table.threshold) == WARNING   # 20.1%, just over
    assert table.worst == WARNING


def test_a_warning_is_not_an_error(tmp_path):
    source = tmp_path / FIRE.name
    shutil.copy2(FIRE, source)
    report = run(source, tmp_path / "out.xlsx", tmp_path / "logs")

    assert report.loss_share_warnings                   # there are warnings
    assert report.loss_share_ok                         # and the run is still sound
    assert report.ok


def test_declared_losses_above_incurred_is_an_error(tmp_path):
    """Individually plausible, jointly impossible — the case R-01 and R-02 miss."""
    nomenclature, blocks = _read(ENGINEERING)
    large = next(b for b in blocks if b.dataset.role == "03")
    for record in large.records:
        record.values["Loss amount"] = 9000.0           # far above any year's incurred

    table = {t.kind: t for t in loss_share_tables(nomenclature, blocks)}["per risk"]
    assert table.worst == EXCEEDS
    assert _row(table, 2023).status(table.threshold) == EXCEEDS


def test_a_year_with_no_losses_shows_zero_not_a_gap():
    """2022 has no Engineering loss of either kind."""
    row = _row(_tables(ENGINEERING)["per risk"], 2022)
    assert row.by_role == {"03": 0.0, "04": 0.0}
    assert row.declared == 0.0
    assert row.incurred == 6350.0
    assert row.status(DEFAULT_THRESHOLD) == OK


# ─────────────────────────────────────────────── what it refuses to do

def test_a_differing_basis_across_the_sections_is_not_summed():
    """Summing a 100% figure with a ceded-only one produces a meaningless number."""
    nomenclature, blocks = _read(FIRE_EQ_WIND)
    wind = next(b for b in blocks if b.dataset.role == "04" and b.section == "Windstorm")
    wind.attributes["Share basis"] = Attribute("Share basis", "ceded only", False, 0)

    table = {t.kind: t for t in loss_share_tables(nomenclature, blocks)}["cat"]
    assert table.skipped
    assert "Share basis differs" in table.skipped
    assert table.rows == []


def test_a_year_only_some_sections_report_is_named_not_assumed():
    """Windstorm drops 2021 from its history; the cat basis for 2021 is then partial."""
    nomenclature, blocks = _read(FIRE_EQ_WIND)
    wind = next(b for b in blocks if b.dataset.role == "01" and b.section == "Windstorm")
    wind.records = [r for r in wind.records if r.values["Year"] != "2021"]

    table = {t.kind: t for t in loss_share_tables(nomenclature, blocks)}["cat"]
    row = _row(table, 2021)
    assert row.partial                                  # EQ reported it, Windstorm did not
    assert row.reporting == 1 and row.expected == 2
    assert row.incurred == 1200.0                       # EQ alone, not 3,500
    assert not _row(table, 2022).partial


def test_the_transposed_twin_is_not_counted_twice():
    """`01. History` and `01. History_Transposed` are the same portfolio."""
    table = _tables(FIRE)["per risk"]
    assert _row(table, 2021).incurred == 14930.0        # not 29,860
    assert table.sheets == ("01. History",)


def test_a_pack_with_no_loss_dataset_gets_no_table():
    nomenclature, blocks = _read(FIRE)
    kept = [b for b in blocks if b.dataset.role in ("01", "02")]
    assert loss_share_tables(nomenclature, kept) == []


# ───────────────────────────────────────────────────── the threshold

class _Nom:
    def __init__(self, value=None):
        self.globals = {} if value is None else {"Loss share warning": value}


@pytest.mark.parametrize("declared,expected", [
    ("20%", 0.20), ("20", 0.20), (20, 0.20), (0.2, 0.2), (0.35, 0.35),
    ("12.5%", 0.125), (None, DEFAULT_THRESHOLD), ("", DEFAULT_THRESHOLD),
    ("nonsense", DEFAULT_THRESHOLD),
])
def test_the_threshold_is_read_from_sheet_00(declared, expected):
    assert threshold_of(_Nom(declared)) == pytest.approx(expected)


def test_the_reference_workbooks_declare_the_threshold():
    for path in (FIRE, ENGINEERING, FIRE_CAT, FIRE_EQ_WIND):
        nomenclature, _ = _read(path)
        assert threshold_of(nomenclature) == pytest.approx(0.20), path.name


def test_a_changed_threshold_changes_the_findings():
    nomenclature, blocks = _read(FIRE)
    nomenclature.globals["Loss share warning"] = "50%"
    table = {t.kind: t for t in loss_share_tables(nomenclature, blocks)}["per risk"]
    assert table.threshold == pytest.approx(0.50)
    assert table.worst == OK                            # nothing reaches 50%


# ──────────────────────────────────────────────────── what gets written

def _text(ws):
    return [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]


def test_the_check_is_written_at_the_end_of_sheet_01(tmp_path):
    source = tmp_path / ENGINEERING.name
    shutil.copy2(ENGINEERING, source)
    out = tmp_path / "out.xlsx"
    run(source, out, tmp_path / "logs")

    wb = load_workbook(out, data_only=True)
    text = _text(wb["01. History"])
    assert any("DECLARED LOSSES AGAINST TOTAL INCURRED" in t for t in text)
    assert "Large losses (03)" in text and "Cat losses (04)" in text
    # It belongs to 01 and nowhere else.
    for sheet in ("03. Large Losses", "04. Cat Losses"):
        assert not any("DECLARED LOSSES" in t for t in _text(wb[sheet]))


def test_it_is_separated_by_three_blank_rows(tmp_path):
    from datatransform.writer import ANCHOR_PREFIX

    source = tmp_path / ENGINEERING.name
    shutil.copy2(ENGINEERING, source)
    out = tmp_path / "out.xlsx"
    run(source, out, tmp_path / "logs")

    ws = load_workbook(out, data_only=True)["01. History"]
    anchor = next(c.row for row in ws.iter_rows() for c in row
                  if isinstance(c.value, str) and c.value.startswith(ANCHOR_PREFIX)
                  and "LOSSSHARE" in c.value)
    for row in range(anchor - 3, anchor):
        assert all(c.value is None for c in ws[row]), f"row {row} is not blank"


def test_a_group_spanning_two_sheets_is_written_on_both(tmp_path):
    source = tmp_path / FIRE_EQ_WIND.name
    shutil.copy2(FIRE_EQ_WIND, source)
    out = tmp_path / "out.xlsx"
    run(source, out, tmp_path / "logs")

    wb = load_workbook(out, data_only=True)
    for sheet in ("01. History EQ", "01. History Wind"):
        text = _text(wb[sheet])
        assert any("DECLARED LOSSES AGAINST TOTAL INCURRED" in t for t in text)
        # and each copy says the other exists, so it reads as one check
        assert any("The same table is written at the end of" in t for t in text)
    # Fire carries its own per-risk table, never the cat one.
    fire = _text(wb["01. History Fire"])
    assert any("Sections summed: Fire" in t for t in fire)
    assert not any("Earthquake, Windstorm" in t for t in fire)


def test_the_process_log_records_every_year(tmp_path):
    source = tmp_path / ENGINEERING.name
    shutil.copy2(ENGINEERING, source)
    logs = tmp_path / "logs"
    run(source, tmp_path / "out.xlsx", logs)

    text = next(logs.glob("*_process.log")).read_text(encoding="utf-8")
    assert "Declared losses against total incurred" in text
    assert "warning above 20% of total incurred" in text
    for year in range(2021, 2026):
        assert f"    {year}   declared" in text
