"""Field resolution and record selection. Specification_v1.md §7, §8."""

from __future__ import annotations

import re

from openpyxl.utils import column_index_from_string as col_idx
from openpyxl.utils import get_column_letter as col_letter

from .coerce import coerce
from .markers import attributes_for, block_indices, structural
from .model import (
    Attribute,
    Block,
    Confidence,
    Dataset,
    ExtractionError,
    FieldType,
    Hypothesis,
    Orientation,
    Record,
)
from .nomenclature import Nomenclature, norm

ERROR_CELLS = {"#REF!", "#N/A", "#VALUE!", "#DIV/0!", "#NAME?", "#NULL!", "#NUM!"}
LEADING_NUMBER = re.compile(r"^\s*(\d+)")


def last_non_empty_row(ws) -> int:
    """``max_row`` counts formatted-but-empty cells, so find the real extent — spec §9.1 O2."""
    for row in range(ws.max_row, 0, -1):
        for col in range(1, ws.max_column + 1):
            if ws.cell(row=row, column=col).value not in (None, ""):
                return row
    return 0


def _is_error(value) -> bool:
    return isinstance(value, str) and value.strip() in ERROR_CELLS


def _check_cell(values_ws, formulas_ws, row: int, col: int, where: str):
    """Spec §7.2 (uncalculated formulas) and §12 T2 (error cells)."""
    v = values_ws.cell(row=row, column=col).value
    if _is_error(v):
        raise ExtractionError(
            f"{where}: error cell {col_letter(col)}{row} contains {v!r} — "
            "error values are fatal and are never coerced"
        )
    if v is None:
        f = formulas_ws.cell(row=row, column=col).value
        if isinstance(f, str) and f.startswith("="):
            raise ExtractionError(
                f"{where}: {col_letter(col)}{row} holds formula {f!r} with no cached value — "
                "workbook not calculated. Open and save it in Excel, then re-run."
            )
    return v


def _lenient(value, block: Block, label: str, ref: str):
    """Coerce an *excluded* record, tolerating failure.

    Excluded records are shown only as an independent control, so a value that cannot
    be typed must not stop a run over records that were never extracted.
    """
    try:
        return coerce(value, block.field_types[label], label, ref, [])
    except ExtractionError:
        return value


def _resolve_attributes(block: Block, markers, nomenclature: Nomenclature) -> None:
    """Attributes, hypotheses and derived confidence — spec §5.

    Three tiers: ⟦GLOBAL⟧ in sheet 00, then sheet-wide markers, then block markers.
    """
    found = attributes_for(markers, block.index)

    for name, value in nomenclature.globals.items():
        if name in block.dataset.attributes and name not in found:
            block.attributes[name] = Attribute(name, value, False, 0)

    for name, marker in found.items():
        if name not in block.dataset.attributes:
            block.hypotheses.append(
                Hypothesis(
                    id="",
                    dataset_key=block.dataset.key,
                    attribute=name,
                    value=marker.value,
                    confidence=Confidence.ASSUMED,
                    source="tool",
                    note="attribute declared in the sheet but not listed in sheet 00",
                )
            )
            continue

        if not nomenclature.validate_value(name, marker.value):
            allowed = " | ".join(nomenclature.vocabulary.get(name, []))
            raise ExtractionError(
                f"{block.sheet_name}: attribute {name!r} has value {marker.value!r}, "
                f"which is not in its vocabulary ({allowed})"
            )

        block.attributes[name] = Attribute(name, marker.value, marker.is_hypothesis, marker.row)

    for name in block.dataset.attributes:
        attr = block.attributes.get(name)
        if attr is None:
            block.hypotheses.append(
                Hypothesis(
                    id="",
                    dataset_key=block.dataset.key,
                    attribute=name,
                    value=None,
                    confidence=Confidence.OPEN,
                    source="tool",
                    note="listed in sheet 00 but never declared in the sheet — nobody decided",
                )
            )
        elif attr.is_hypothesis:
            block.hypotheses.append(
                Hypothesis(
                    id="",
                    dataset_key=block.dataset.key,
                    attribute=name,
                    value=attr.value,
                    confidence=Confidence.ASSUMED,
                    source="H_ marker",
                    note=f"declared as a hypothesis at A{attr.row}",
                )
            )


def _occurrence_date_field(block: Block) -> str | None:
    """The one date that defines the occurrence year — spec §9.2.1 S15, §10.2.

    A cat event carries two dates, and only the one the cedent counts from can be
    checked against the year. An event running 2024-12-30 → 2025-01-02 is not
    misfiled because its *end* falls in the following year; comparing every date
    field would report that as an error.
    """
    dates = [h for h in block.fields if block.field_types.get(h) is FieldType.DATE]
    if not dates:
        return None
    declared = block.attributes.get("Occurrence year from")
    if declared is not None:
        named = norm(declared.value)
        if named not in dates:
            raise ExtractionError(
                f"{block.sheet_name!r} block {block.index}: "
                f"'Occurrence year from' = {named!r}, which is not a date field of "
                f"this block ({', '.join(dates)})"
            )
        return named
    return dates[0]


def _check_occurrence_year(block: Block) -> None:
    """Under an occurrence basis the year must match the loss date — spec §10.2.

    Under an underwriting basis it legitimately need not: a policy incepting in one
    year can produce a loss in the next. So the check runs only where the declared
    basis makes it meaningful, and stays silent otherwise.
    """
    basis = block.attributes.get("Year basis")
    if basis is None or norm(basis.value).casefold() != "occurrence":
        return

    key = block.dataset.key_field
    field_name = _occurrence_date_field(block)
    if field_name is None:
        return

    mismatched = []
    for record in block.records:
        label = record.values.get(key)
        match = LEADING_NUMBER.match(label) if isinstance(label, str) else None
        if match is None:
            continue
        value = record.values.get(field_name)
        if value is not None and value.year != int(match.group(1)):
            mismatched.append(f"{record.source_ref}: {label} vs {value.isoformat()}")

    if mismatched:
        block.hypotheses.append(
            Hypothesis(
                id="",
                dataset_key=block.dataset.key,
                attribute="Occurrence year mismatch",
                value=f"{len(mismatched)} record(s)",
                confidence=Confidence.OPEN,
                source="tool",
                note=("Year basis is Occurrence, so the year must equal the year of the "
                      "loss date: " + "; ".join(mismatched[:5])),
            )
        )


def _check_period_overlap(block: Block) -> None:
    """Repeated leading year means overlapping periods — spec §9.2 S10.

    ``2026`` and ``2026 9 months`` are two records of the same year, so a column total
    counts that year twice. The tool flags it and changes nothing: which record to use
    is an underwriting judgment, not the tool's to make.

    Only *distinct* labels count. Eight large losses sharing the year 2021 are eight
    claims, not an overlapping period — a transactional listing has many records per
    year by nature, and flagging that would bury the real signal.
    """
    key = block.dataset.key_field
    if block.field_types.get(key) is not FieldType.TEXT:
        return

    groups: dict[str, list[str]] = {}
    for record in block.records:
        label = record.values.get(key)
        if not isinstance(label, str):
            continue
        match = LEADING_NUMBER.match(label)
        if match and label not in groups.setdefault(match.group(1), []):
            groups[match.group(1)].append(label)

    for year, labels in sorted(groups.items()):
        if len(labels) > 1:
            block.hypotheses.append(
                Hypothesis(
                    id="",
                    dataset_key=block.dataset.key,
                    attribute=f"Period overlap {year}",
                    value=" · ".join(labels),
                    confidence=Confidence.ASSUMED,
                    source="tool",
                    note=(
                        f"{key} {year} appears {len(labels)} times as overlapping "
                        "periods; the control sum counts each of them, so it verifies "
                        "extraction rather than being a portfolio total"
                    ),
                )
            )


def _number_hypotheses(block: Block) -> None:
    for n, hypothesis in enumerate(block.hypotheses, start=1):
        hypothesis.id = f"H-{n:02d}"


def _resolve_labels(cells: list[tuple], dataset: Dataset, where: str) -> dict[str, str]:
    """Strict label match, duplicates and absences fatal — spec §8 F1–F4.

    Takes a list of ``(text, ref)`` pairs rather than a mapping: keying by cell value
    would collapse duplicates, which are precisely what F3 must detect.
    """
    wanted = {norm(h).casefold(): h for h in dataset.headers}
    hits: dict[str, list[int]] = {}
    for text, ref in cells:
        key = norm(text).casefold()
        if key in wanted:
            hits.setdefault(wanted[key], []).append(ref)

    for label, refs in hits.items():
        if len(refs) > 1:
            raise ExtractionError(
                f"{where}: label {label!r} appears {len(refs)} times in the extraction "
                f"row/column — ambiguous"
            )
    missing = [h for h in dataset.headers if h not in hits and not dataset.is_optional(h)]
    if missing:
        raise ExtractionError(
            f"{where}: label(s) {missing} declared in sheet 00 but absent from the "
            f"extraction row/column"
        )
    return {label: str(refs[0]) for label, refs in hits.items()}


def _block_boundary(markers, index: int, orientation, limit_row: int) -> int:
    """Where this block's records stop — spec §6, §6.5.

    Two arrangements, and the selector column tells them apart without anyone declaring
    which is which:

    **Stacked** — blocks sit one below the other and *share* a selector column, as the
    section-blocks of a combined History sheet do. Without a boundary the first block
    would scan to the end and swallow the records below it.

    **Overlaid** — blocks read the *same* rows through *different* selector columns, as
    a combined large-and-cat loss list does. Here a boundary would be fatal: the second
    block's extraction row sits above the shared records, so the first block would stop
    before reading any of them.

    A block therefore stops only at a later block that could compete for its rows — one
    reading the same selector column. Where the columns differ, each selector already
    identifies its own rows and no boundary is needed.
    """
    if orientation is Orientation.TRANSPOSED:
        return limit_row
    mine = next((m.row for m in markers if m.name == "Header" and m.index == index), None)
    if mine is None:
        return limit_row

    my_selector = _selector_of(markers, index)
    later = [
        m.row for m in markers
        if m.name == "Header" and m.index != index and m.row > mine
        and _selector_of(markers, m.index) == my_selector
    ]
    return min(later) - 1 if later else limit_row


def _selector_of(markers, index: int) -> str | None:
    marker = structural(markers, "Info", index)
    return norm(marker.value).casefold() if marker and marker.value else None


def _header_rows(markers) -> set[int]:
    """Every block's extraction row. None of them is ever a record — spec §6.3.

    Matters for overlaid blocks: block 1's records begin one row below its own header
    and would otherwise swallow block 2's extraction row, whose labels sit in the same
    columns block 1 reads as measures.
    """
    return {m.row for m in markers if m.name == "Header"}


def extract_block(values_ws, formulas_ws, dataset: Dataset, markers, index: int,
                  nomenclature: Nomenclature) -> Block:
    header = structural(markers, "Header", index)
    info = structural(markers, "Info", index)
    transposed = structural(markers, "Transpose", index) is not None
    where = f"{values_ws.title} block {index}"

    if info is None or not info.value:
        raise ExtractionError(f"{where}: Header_{index} present but Info_{index} missing")

    orientation = Orientation.TRANSPOSED if transposed else Orientation.ROW_WISE
    limit_row = last_non_empty_row(values_ws)
    limit_col = values_ws.max_column

    if orientation is Orientation.ROW_WISE:
        if header.value:
            raise ExtractionError(
                f"{where}: Header_{index} carries a value but Transpose_{index} is absent"
            )
        header_ref, info_ref = str(header.row), info.value.strip().upper()
        cells = [
            (values_ws.cell(row=header.row, column=c).value, c)
            for c in range(1, limit_col + 1)
            if values_ws.cell(row=header.row, column=c).value is not None
        ]
        address = _resolve_labels(cells, dataset, where)
        address = {label: col_letter(int(ref)) for label, ref in address.items()}
    else:
        if not header.value:
            raise ExtractionError(
                f"{where}: Transpose_{index} present, so Header_{index} must name a column"
            )
        header_ref, info_ref = header.value.strip().upper(), info.value.strip()
        hcol = col_idx(header_ref)
        cells = [
            (values_ws.cell(row=r, column=hcol).value, r)
            for r in range(1, limit_row + 1)
            if values_ws.cell(row=r, column=hcol).value is not None
        ]
        address = _resolve_labels(cells, dataset, where)

    block = Block(
        dataset=dataset,
        sheet_name=values_ws.title,
        index=index,
        orientation=orientation,
        header_ref=header_ref,
        info_ref=info_ref,
        address_map=address,
        field_types=nomenclature.field_types(dataset.headers),
    )
    block.section = None
    _resolve_attributes(block, markers, nomenclature)
    section = block.attributes.get("Section")
    block.section = norm(section.value) if section else None

    if orientation is Orientation.ROW_WISE:
        boundary = _block_boundary(markers, index, orientation, limit_row)
        _select_rows(block, values_ws, formulas_ws, boundary, limit_col, where,
                     _header_rows(markers))
    else:
        _select_columns(block, values_ws, formulas_ws, limit_row, limit_col, where)

    _check_period_overlap(block)
    _check_occurrence_year(block)
    _number_hypotheses(block)
    return block


def _select_rows(block: Block, values_ws, formulas_ws, limit_row, limit_col, where,
                 header_rows=()):
    sel = col_idx(block.info_ref)
    header_row = int(block.header_ref)
    # A candidate must carry at least one measure; text in the key column (a footnote,
    # a section caption) is not a record that was excluded.
    measure_cols = [col_idx(block.address_map[m]) for m in block.numeric_fields
                    if m in block.address_map]

    # Records follow the extraction row; it is never itself a record — spec §6.3, §7.
    for row in range(header_row + 1, limit_row + 1):
        if row in header_rows:
            continue                      # another block's extraction row, not a record
        marker = values_ws.cell(row=row, column=sel).value
        # A selector holding "" is a formula that decided this row is not ours — spec
        # §7.1. `=IF(<test>, ROW(), "")` is how one list is split into two blocks.
        if isinstance(marker, str) and not marker.strip():
            marker = None
        if marker is None:
            formula = formulas_ws.cell(row=row, column=sel).value
            if isinstance(formula, str) and formula.startswith("="):
                raise ExtractionError(
                    f"{where}: selector {block.info_ref}{row} holds {formula!r} with no cached "
                    "value — workbook not calculated. A skipped record would be silent."
                )
            if any(values_ws.cell(row=row, column=c).value is not None
                   for c in measure_cols):
                block.candidates += 1
                block.excluded += 1
                block.excluded_records.append(
                    Record(
                        source_ref=str(row),
                        values={
                            label: _lenient(values_ws.cell(row=row, column=col_idx(col)).value,
                                            block, label, f"{col}{row}")
                            for label, col in block.address_map.items()
                        },
                    )
                )
            continue

        block.candidates += 1
        if marker != row:
            raise ExtractionError(
                f"{where}: selector {block.info_ref}{row} reads {marker!r} but sits on row "
                f"{row} — the sheet has been manipulated and the markers no longer align"
            )
        values = {}
        date_format = block.attributes.get("Date format")
        for label, col in block.address_map.items():
            raw = _check_cell(values_ws, formulas_ws, row, col_idx(col), where)
            values[label] = coerce(raw, block.field_types[label], label,
                                   f"{col}{row}", block.coercions,
                                   date_format.value if date_format else None)
        block.records.append(Record(source_ref=str(row), values=values))

    extracted = {col_idx(c) for c in block.address_map.values()}
    rows = [int(r.source_ref) for r in block.records]
    for col in range(2, limit_col + 1):
        if col == sel or col in extracted:
            continue
        if any(values_ws.cell(row=r, column=col).value is not None for r in rows):
            block.unextracted.append(col_letter(col))


def _select_columns(block: Block, values_ws, formulas_ws, limit_row, limit_col, where):
    sel = int(block.info_ref)
    header_col = col_idx(block.header_ref)
    measure_rows = [int(block.address_map[m]) for m in block.numeric_fields
                    if m in block.address_map]

    # Records follow the extraction column; column A is the marker channel — spec §4 M1, §6.3.
    for col in range(header_col + 1, limit_col + 1):
        marker = values_ws.cell(row=sel, column=col).value
        if marker is None:
            formula = formulas_ws.cell(row=sel, column=col).value
            if isinstance(formula, str) and formula.startswith("="):
                raise ExtractionError(
                    f"{where}: selector {col_letter(col)}{sel} holds {formula!r} with no cached "
                    "value — workbook not calculated. A skipped record would be silent."
                )
            if any(values_ws.cell(row=r, column=col).value is not None
                   for r in measure_rows):
                block.candidates += 1
                block.excluded += 1
                block.excluded_records.append(
                    Record(
                        source_ref=col_letter(col),
                        values={
                            label: _lenient(values_ws.cell(row=int(row), column=col).value,
                                            block, label, f"{col_letter(col)}{row}")
                            for label, row in block.address_map.items()
                        },
                    )
                )
            continue

        block.candidates += 1
        if marker != col:
            raise ExtractionError(
                f"{where}: selector {col_letter(col)}{sel} reads {marker!r} but sits in column "
                f"{col_letter(col)} ({col}) — the sheet has been manipulated"
            )
        values = {}
        date_format = block.attributes.get("Date format")
        for label, row in block.address_map.items():
            raw = _check_cell(values_ws, formulas_ws, int(row), col, where)
            values[label] = coerce(raw, block.field_types[label], label,
                                   f"{col_letter(col)}{row}", block.coercions,
                                   date_format.value if date_format else None)
        # Row-level confidence applies to transposed blocks only — spec §5.3, user ad 7.
        block.records.append(
            Record(source_ref=col_letter(col), values=values, confidence=block.confidence)
        )

    extracted = {int(r) for r in block.address_map.values()}
    cols = [col_idx(r.source_ref) for r in block.records]
    for row in range(1, limit_row + 1):
        if row == sel or row in extracted:
            continue
        if any(values_ws.cell(row=row, column=c).value is not None for c in cols):
            block.unextracted.append(str(row))


def _dataset_for_block(values_ws, nomenclature, markers, index, default):
    """Which dataset this block *is* — spec §6.4.

    The sheet name is a **default, not a law**. A cedent who reports large losses and
    cat events in one list needs both datasets on one sheet, so a block may name its
    own with ``Dataset_i = 04 Cat``.
    """
    declared = structural(markers, "Dataset", index)
    if declared is None or not declared.value:
        if default is None:
            raise ExtractionError(
                f"{values_ws.title!r} block {index}: the sheet matches no dataset in "
                f"sheet 00, so the block must name one with 'Dataset_{index} = <key>'"
            )
        return default

    wanted = norm(declared.value)
    for dataset in nomenclature.datasets.values():
        if norm(dataset.key).casefold() == wanted.casefold():
            return dataset
    raise ExtractionError(
        f"{values_ws.title!r} block {index}: Dataset_{index} = {wanted!r}, which is not "
        f"a dataset key in sheet 00 ({', '.join(d.key for d in nomenclature.datasets.values())})"
    )


def extract_sheet(values_ws, formulas_ws, nomenclature: Nomenclature, markers) -> list[Block]:
    """All blocks of one sheet. No ``Header_i`` → nothing extracted — spec §4 M7."""
    default = nomenclature.dataset_for(values_ws.title)
    indices = block_indices(markers)
    if default is None and not any(
        structural(markers, "Dataset", i) is not None for i in indices
    ):
        return []

    blocks = [
        extract_block(
            values_ws, formulas_ws,
            _dataset_for_block(values_ws, nomenclature, markers, i, default),
            markers, i, nomenclature,
        )
        for i in indices
    ]
    _reconcile_siblings(blocks)
    return blocks


def _reconcile_siblings(blocks) -> None:
    """Reconcile blocks that overlay one another — spec §6.5.

    Where two blocks read the same rows through different selectors, each sees the
    other's rows and columns and would otherwise report them as findings against
    itself. Both statements would be false in an audit document:

    * a column the sibling extracts is **not** "contained data, not declared in 00" —
      it is declared, in the other block's dataset;
    * a row the sibling extracts is **not** an excluded record — nothing was dropped,
      it simply belongs to the other list.
    """
    taken = {ref for b in blocks for ref in b.address_map.values()}
    for block in blocks:
        block.unextracted = [ref for ref in block.unextracted if ref not in taken]

    for block in blocks:
        elsewhere = {
            r.source_ref: b.dataset.key
            for b in blocks if b is not block
            for r in b.records
        }
        kept = [r for r in block.excluded_records if r.source_ref not in elsewhere]
        claimed = len(block.excluded_records) - len(kept)
        if claimed:
            block.excluded_records = kept
            block.excluded -= claimed
            block.claimed_elsewhere = claimed
            block.claimed_by = sorted(set(elsewhere.values()))
