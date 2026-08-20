"""05. Risk Profiles — specification_v1.md §2.4, §9.2.1 S16/S17."""

from __future__ import annotations

import shutil
from pathlib import Path

import pytest
from openpyxl import load_workbook

from datatransform.bands import BandError, continuity, parse_band
from datatransform.model import ExtractionError
from datatransform.extract import extract_sheet, last_non_empty_row
from datatransform.markers import read_markers
from datatransform.nomenclature import read_nomenclature
from datatransform.runner import run
from datatransform.specs import step2_for
from datatransform.transform import apply_step2

ROOT = Path(__file__).resolve().parents[1]

FIRE = ROOT / "Intake_v1.xlsx"                      # bounds declared as columns
ENGINEERING = ROOT / "Intake_Engineering_v1.xlsx"   # bounds read off the label
COMBINED = ROOT / "Intake_EngineeringCombined_v1.xlsx"
FIRE_CAT = ROOT / "Intake_FireCat_v1.xlsx"
FIRE_EQ_WIND = ROOT / "Intake_FireEQWind_v1.xlsx"


def _read(path):
    values = load_workbook(path, data_only=True)
    formulas = load_workbook(path, data_only=False)
    nomenclature = read_nomenclature(formulas)
    blocks = []
    for title in formulas.sheetnames:
        if title.startswith("00."):
            continue
        markers = read_markers(formulas[title], last_non_empty_row(values[title]))
        blocks.extend(extract_sheet(values[title], formulas[title], nomenclature, markers))
    return nomenclature, blocks


def _profiles(path):
    nomenclature, blocks = _read(path)
    return nomenclature, [b for b in blocks if b.dataset.role == "05"]


def _step2(path, section=None):
    nomenclature, profiles = _profiles(path)
    block = next(b for b in profiles if section is None or b.section == section)
    return block, apply_step2(block, step2_for(block.dataset.key), nomenclature)


# ─────────────────────────────────────── one profile per line of business

@pytest.mark.parametrize("path,sections", [
    (FIRE, {"Fire"}),
    (ENGINEERING, {"Engineering"}),
    (COMBINED, {"Engineering"}),
    (FIRE_CAT, {"Fire", "Earthquake", "Windstorm"}),
    (FIRE_EQ_WIND, {"Fire", "Earthquake", "Windstorm"}),
])
def test_every_line_of_business_has_a_profile(path, sections):
    nomenclature, profiles = _profiles(path)
    assert {b.section for b in profiles} == sections
    # and every section declares that it expects one
    for section in nomenclature.sections:
        assert section.expects("05"), section.name


def test_the_two_cat_shapes_produce_the_same_profiles():
    """One sheet with three blocks, or three sheets — the same profiles."""
    _, combined = _profiles(FIRE_CAT)
    _, separate = _profiles(FIRE_EQ_WIND)

    def shape(blocks):
        return {b.section: [r.values for r in b.records] for b in blocks}

    assert shape(combined) == shape(separate)


# ────────────────────────────────────────────────── bounds: two shapes

def test_declared_bounds_are_extracted_not_derived():
    block, result = _step2(FIRE)
    assert "Band from" in block.fields
    assert result.derived_fields == ()
    assert not any("read off" in n for n in result.notes)


def test_absent_bounds_are_read_off_the_label_in_step_2():
    block, result = _step2(ENGINEERING)
    assert "Band from" not in block.fields          # step 1 has only what the sheet says
    assert result.derived_fields == ("Band from", "Band to")

    assert [r.values["Band from"] for r in result.records] == [0, 1001, 5001, 20000]
    assert [r.values["Band to"] for r in result.records] == [1000, 5000, 20000, None]
    assert any("read off Band" in n for n in result.notes)


def test_step_1_never_gains_the_derived_bounds():
    """S11: step 1 carries only extracted information."""
    block, result = _step2(ENGINEERING)
    assert all("Band from" not in r.values for r in block.records)
    assert all("Band from" in r.values for r in result.records)


def test_the_derived_bounds_appear_in_the_declared_order():
    _, result = _step2(ENGINEERING)
    assert result.declared_columns == (
        "Band", "Band from", "Band to", "Premium", "Number of Risks", "Exposure",
    )


# ─────────────────────────────────────────────────────── the parser

@pytest.mark.parametrize("label,expected", [
    ("0 - 1 000", (0.0, 1000.0)),
    ("1 001 – 5 000", (1001.0, 5000.0)),
    ("5 001 to 20 000", (5001.0, 20000.0)),
    ("> 25 000", (25000.0, None)),
    (">= 25 000", (25000.0, None)),
    ("over 1 000 000", (1000000.0, None)),
    ("< 10 000", (None, 10000.0)),
    ("up to 10 000", (None, 10000.0)),
    ("1 Mio - 5 Mio", (1_000_000.0, 5_000_000.0)),
    ("10k - 50k", (10_000.0, 50_000.0)),
    ("2 500", (2500.0, 2500.0)),
])
def test_unambiguous_labels_are_read(label, expected):
    assert parse_band(label, "B7") == expected


@pytest.mark.parametrize("label", ["", "  ", "band", "big risks", "1 000 - "])
def test_a_label_that_cannot_be_read_is_fatal(label):
    with pytest.raises(BandError):
        parse_band(label, "B7")


def test_an_ambiguous_separator_is_refused_and_says_how_to_fix_it():
    """1,000 could be one thousand or one point zero — the decimal rule, §8.4."""
    with pytest.raises(BandError, match="ambiguous"):
        parse_band("0 - 1,000", "B7")
    with pytest.raises(BandError, match="Declare 'Band from' and 'Band to' as columns"):
        parse_band("0 - 1,000", "B7")


def test_an_upper_bound_below_the_lower_is_refused():
    with pytest.raises(BandError, match="below the lower bound"):
        parse_band("10 000 - 5 000", "B7")


def test_an_open_bound_is_absent_not_zero():
    assert parse_band("> 25 000", "B7") == (25000.0, None)
    assert parse_band("< 25 000", "B7") == (None, 25000.0)


# ──────────────────────────────────────────────────── continuity

def test_both_band_conventions_pass():
    shared = [("a", 0, 10_000), ("b", 10_000, 20_000)]
    gapless = [("a", 1, 10_000), ("b", 10_001, 20_000)]
    assert continuity(shared) == []
    assert continuity(gapless) == []


def test_a_real_gap_is_reported():
    findings = continuity([("0 – 10 000", 0, 10_000), ("20 001 – 30 000", 20_001, 30_000)])
    assert len(findings) == 1
    assert "in no band" in findings[0]


def test_an_overlap_is_reported():
    findings = continuity([("0 – 10 000", 0, 10_000), ("5 000 – 20 000", 5_000, 20_000)])
    assert len(findings) == 1
    assert "overlap" in findings[0]


def test_an_open_bound_ends_the_chain():
    assert continuity([("a", 0, None), ("b", 50_000, None)]) == []


@pytest.mark.parametrize("path", [FIRE, ENGINEERING, FIRE_CAT, FIRE_EQ_WIND])
def test_the_reference_profiles_are_continuous(path):
    nomenclature, profiles = _profiles(path)
    for block in profiles:
        result = apply_step2(block, step2_for(block.dataset.key), nomenclature)
        assert not [n for n in result.notes if n.startswith("BAND CONTINUITY")], \
            f"{path.name} / {block.section}"


# ──────────────────────────────────────────────────── step 2 output

def test_sorted_by_the_lower_bound_as_a_number():
    """Alphanumeric would put '10 001 – 25 000' before '5 001 – 10 000'."""
    _, result = _step2(FIRE)
    lower = [r.values["Band from"] for r in result.records]
    assert lower == sorted(lower)
    assert [r.values["Band"] for r in result.records][:3] == [
        "0 - 1 000", "1 001 - 5 000", "5 001 - 10 000",
    ]


def test_cumulative_shares_run_to_one():
    _, result = _step2(FIRE)
    for name in ("Cumulative risks %", "Cumulative exposure %", "Cumulative premium %"):
        column = result.computed[name]
        assert column == sorted(column), f"{name} must be non-decreasing"
        assert column[-1] == pytest.approx(1.0)


def test_cumulative_premium_answers_where_the_book_sits():
    _, result = _step2(FIRE)
    # 2,150 + 4,900 of 20,150 sits at or below 5,000
    assert result.computed["Cumulative premium %"][1] == pytest.approx(7050 / 20150)


def test_per_band_measures_are_computed():
    _, result = _step2(FIRE)
    # band 1: 310,000 exposure over 620 risks; 2,150 premium on 310,000 exposure
    assert result.computed["Average exposure per risk"][0] == pytest.approx(500.0)
    assert result.computed["Rate on exposure ‰"][0] == pytest.approx(2150 / 310000 * 1000)


def test_bounds_are_never_summed():
    """A bound is a number, not a quantity — spec §2.4."""
    block, result = _step2(FIRE)
    assert "Band from" in block.numeric_fields
    assert "Band from" not in block.measure_fields
    assert set(result.totals()) == {"Premium", "Number of Risks", "Exposure"}
    assert result.totals()["Premium"] == 20150.0


def test_step_2_is_value_preserving():
    block, result = _step2(FIRE)
    assert block.totals() == result.totals()


# ────────────────── one profile, three presentations — spec §2.4

def _demo(tmp_path):
    """The three-shape demonstration workbook, built fresh."""
    import sys

    sys.path.insert(0, str(ROOT / "tools"))
    from demo_profile_shapes import build

    path, _, _ = build(tmp_path / "shapes.xlsx")
    nomenclature, blocks = _read(path)
    return path, {b.sheet_name: (b, apply_step2(b, step2_for(b.dataset.key), nomenclature))
                  for b in blocks}


def test_every_presentation_of_one_profile_agrees(tmp_path):
    """Two columns, one label, bounds without a label, European grouping — one profile."""
    _, shapes = _demo(tmp_path)
    assert set(shapes) == {"05. Profile two columns", "05. Profile one column",
                           "05. Profile bounds only", "05. Profile European"}

    def figures(result):
        return [(r.values["Band from"], r.values["Band to"], r.values["Premium"],
                 r.values["Number of Risks"], r.values["Exposure"])
                for r in result.records]

    reference = figures(shapes["05. Profile two columns"][1])
    assert reference[0][:2] == (1.0, 1_000_000.0)          # the bounds you named
    assert reference[-1][:2] == (25_000_000.0, None)       # open at the top
    for name, (_, result) in shapes.items():
        assert figures(result) == reference, name
        assert result.totals()["Premium"] == 8_000_000.0


def test_only_the_declared_shape_extracts_the_bounds_in_step_1(tmp_path):
    _, shapes = _demo(tmp_path)
    two_columns = shapes["05. Profile two columns"]
    one_column = shapes["05. Profile one column"]

    assert "Band from" in two_columns[0].fields
    assert two_columns[1].derived_fields == ()

    assert "Band from" not in one_column[0].fields
    assert one_column[1].derived_fields == ("Band from", "Band to")


def test_step_2_always_returns_both_bounds(tmp_path):
    """However the source presented them — spec §2.4."""
    _, shapes = _demo(tmp_path)
    for name, (_, result) in shapes.items():
        assert "Band from" in result.declared_columns, name
        assert "Band to" in result.declared_columns, name
        assert all(r.values["Band from"] is not None for r in result.records), name


def test_a_profile_needs_no_label_column(tmp_path):
    """The bounds carry the band's identity; the label is a human convenience."""
    _, shapes = _demo(tmp_path)
    block, result = shapes["05. Profile bounds only"]

    assert "Band" not in block.fields               # never declared, never read
    assert result.derived_fields == ()              # nothing to work out
    assert result.declared_columns == (
        "Band from", "Band to", "Premium", "Number of Risks", "Exposure",
    )
    assert [r.values["Band from"] for r in result.records] == [
        1, 1_000_001, 5_000_001, 25_000_000,
    ]


def test_a_block_with_neither_label_nor_bounds_is_fatal():
    """Nothing to extract and nothing to read off — say so rather than sort on nothing."""
    from datatransform.specs import Bounds, Step2Spec

    nomenclature, profiles = _profiles(FIRE)
    block = profiles[0]
    block.address_map = {k: v for k, v in block.address_map.items()
                         if k not in ("Band", "Band from", "Band to")}
    spec = Step2Spec(sort_by=("Premium",),
                     bounds=Bounds("Band", "Band from", "Band to"))
    with pytest.raises(ExtractionError, match="band bounds cannot be produced"):
        apply_step2(block, spec, nomenclature)


def test_both_thousands_conventions_are_read(tmp_path):
    """1-1,000,000 and 1-1.000.000 both group unambiguously, so neither is guessed."""
    _, shapes = _demo(tmp_path)
    anglo = shapes["05. Profile one column"][1]
    european = shapes["05. Profile European"][1]

    assert anglo.records[0].values["Band"] == "1-1,000,000"
    assert european.records[0].values["Band"] == "1-1.000.000"
    assert [r.values["Band from"] for r in anglo.records] == \
           [r.values["Band from"] for r in european.records]


# ───────────────────────────────────────────────────────── end to end

@pytest.mark.parametrize("path", [FIRE, ENGINEERING, COMBINED, FIRE_CAT, FIRE_EQ_WIND])
def test_profiles_run_end_to_end(path, tmp_path):
    source = tmp_path / path.name
    shutil.copy2(path, source)
    report = run(source, tmp_path / "out.xlsx", tmp_path / "logs")
    assert report.ok, [o.detail for o in report.outcomes if o.status == "error"]
    assert any(o.sheet.startswith("05.") and o.status == "processed"
               for o in report.outcomes)


def test_the_written_profile_carries_its_derived_columns(tmp_path):
    source = tmp_path / ENGINEERING.name
    shutil.copy2(ENGINEERING, source)
    out = tmp_path / "out.xlsx"
    run(source, out, tmp_path / "logs")

    ws = load_workbook(out, data_only=True)["05. Risk Profiles"]
    text = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
    for column in ("Band from", "Band to", "Average exposure per risk",
                   "Rate on exposure ‰", "Cumulative premium %"):
        assert column in text
    assert any("read off Band" in t for t in text)
    # Step 1 says the columns were absent; step 2 says it worked them out.
    assert any("absent from this block: Band from, Band to" in t for t in text)
