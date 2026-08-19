"""Writing step blocks beneath the original data. Specification_v1.md §9, §10."""

from __future__ import annotations

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as col_letter

from .model import Block, Confidence
from .specs import SPEC_VERSION
from .transform import Step2Result

FONT = "Arial"
ANCHOR_PREFIX = "⟦DT:"
GAP = 3                       # blank rows between blocks — spec §9.1 O2/O3
FIRST_COL = 2                 # column B; column A stays the marker channel

anchor_f = Font(name=FONT, size=9, bold=True, color="7F6000")
title_f = Font(name=FONT, size=11, bold=True)
note_f = Font(name=FONT, size=9, italic=True, color="595959")
head_f = Font(name=FONT, size=10, bold=True)
body_f = Font(name=FONT, size=10)
ctrl_f = Font(name=FONT, size=10, bold=True)

step1_fill = PatternFill("solid", fgColor="DDEBF7")
step2_fill = PatternFill("solid", fgColor="E2EFDA")
ctrl_fill = PatternFill("solid", fgColor="FFF2CC")


def anchor_tag(key: str, step: str) -> str:
    return f"{ANCHOR_PREFIX}{key}:{step}:v{SPEC_VERSION}⟧"


def clear_generated(ws) -> int:
    """Re-runs replace rather than stack — spec §9.1 O5."""
    anchors = [
        r for r in range(1, ws.max_row + 1)
        if isinstance(ws.cell(row=r, column=FIRST_COL).value, str)
        and ws.cell(row=r, column=FIRST_COL).value.startswith(ANCHOR_PREFIX)
    ]
    if not anchors:
        return 0
    start = max(1, min(anchors) - GAP)
    removed = ws.max_row - start + 1
    if removed > 0:
        ws.delete_rows(start, removed)
    return removed


class BlockWriter:
    """Writes the two step blocks and reports the cached values its formulas need."""

    def __init__(self, ws, start_row: int):
        self.ws = ws
        self.row = start_row
        self.formula_values: list[tuple[str, str, float]] = []

    def _put(self, col: int, value, font=body_f, fmt=None, fill=None):
        cell = self.ws.cell(row=self.row, column=col, value=value)
        cell.font = font
        if fmt:
            cell.number_format = fmt
        if fill:
            cell.fill = fill
        return cell

    def _line(self, text, font=note_f, col=FIRST_COL):
        self._put(col, text, font)
        self.row += 1

    def _formula(self, col: int, formula: str, expected, fmt=None, font=ctrl_f, fill=None):
        cell = self._put(col, formula, font, fmt, fill)
        if expected is not None:
            self.formula_values.append((self.ws.title, cell.coordinate, expected))
        return cell

    # ────────────────────────────────────────────────────────── step 1
    def write_step1(self, block: Block) -> None:
        ds = block.dataset
        self._put(FIRST_COL, anchor_tag(ds.key, "STEP1"), anchor_f)
        self.row += 1
        self._put(FIRST_COL, "STEP 1 — FILTERED & INTERPRETED", title_f, fill=step1_fill)
        self.row += 1

        self._line(
            f"Source: {block.sheet_name} · block {block.index} · {block.orientation.value} · "
            f"extraction {'row' if block.orientation.value == 'row-wise' else 'column'} "
            f"{block.header_ref} · selector {block.info_ref}"
        )
        self._line(
            "Fields resolved by label: "
            + " · ".join(f"{k} → {v}" for k, v in block.address_map.items())
        )
        absent = [h for h in block.dataset.headers if h not in block.address_map]
        if absent:
            # Declared optional in sheet 00 and not supplied here. Saying so is the
            # difference between "the cedent did not report it" and "the tool lost it".
            self._line(
                "Declared optional in sheet 00, absent from this block: "
                + ", ".join(absent)
            )
        declared = " · ".join(
            f"{a.name} = {a.value}{' (hypothesis)' if a.is_hypothesis else ''}"
            for a in block.attributes.values()
        )
        self._line(f"Attributes: {declared}" if declared else "Attributes: none declared")
        counted = (
            f"Records: {block.candidates} candidates, {len(block.records)} extracted, "
            f"{block.excluded} excluded (selector empty)"
        )
        if block.claimed_elsewhere:
            counted += (
                f", {block.claimed_elsewhere} extracted by "
                f"{' and '.join(block.claimed_by)} on this sheet — one list, two datasets"
            )
        self._line(counted)
        self._line(
            "Types applied: "
            + " · ".join(f"{h} → {block.field_types[h].value}" for h in block.fields)
        )
        notable = [c for c in block.coercions if not c.routine]
        if notable:
            shown = "; ".join(
                f"{c.field}@{c.source_ref} {c.before!r}→{c.after!r} ({c.note})"
                for c in notable[:6]
            )
            more = f" (+{len(notable) - 6} more)" if len(notable) > 6 else ""
            self._line(f"Values that needed conversion: {shown}{more}")
        self._line(f"Confidence: {block.confidence.value} (worst of all hypotheses)")
        if block.unextracted:
            self._line(
                "Contained data but not extracted, not declared in sheet 00: "
                + ", ".join(block.unextracted)
            )
        self.row += 1

        if block.hypotheses:
            self._put(FIRST_COL, "Assumptions & hypotheses", head_f)
            self.row += 1
            for h in block.hypotheses:
                self._put(FIRST_COL, h.id, body_f)
                self._put(FIRST_COL + 1, h.attribute, body_f)
                self._put(FIRST_COL + 2, "—" if h.value is None else str(h.value), body_f)
                self._put(FIRST_COL + 3, h.confidence.value, body_f)
                self._put(FIRST_COL + 4, h.status, body_f)
                self._put(FIRST_COL + 5, h.note, note_f)
                self.row += 1
            self.row += 1

        columns = [block.provenance_label, *block.fields]
        if block.orientation.value == "transposed":
            columns.append("Confidence")

        for i, label in enumerate(columns):
            self._put(FIRST_COL + i, label, head_f, fill=step1_fill)
        self.row += 1

        first_data = self.row
        for record in block.records:
            self._put(FIRST_COL, record.source_ref, body_f)
            for i, label in enumerate(block.fields, start=1):
                self._put(FIRST_COL + i, record.values.get(label), body_f,
                          block.number_format(label))
            if block.orientation.value == "transposed":
                self._put(FIRST_COL + len(block.fields) + 1,
                          record.confidence.value if record.confidence else "", note_f)
            self.row += 1
        last_data = self.row - 1

        self._write_controls(block, block.fields, first_data, last_data, block.totals())
        self._write_excluded_note(block)

    def _write_controls(self, block, headers, first_data, last_data, totals):
        self._put(FIRST_COL, f"Control  (n = {len(block.records)})", ctrl_f, fill=ctrl_fill)
        control_row = self.row
        for i, label in enumerate(headers, start=1):
            if label not in block.measure_fields:
                continue
            letter = col_letter(FIRST_COL + i)
            self._formula(FIRST_COL + i, f"=SUM({letter}{first_data}:{letter}{last_data})",
                          totals.get(label), block.number_format(label), ctrl_f, ctrl_fill)
        self.row += 1

        self._put(FIRST_COL, "Expected (computed by tool)", note_f)
        expected_row = self.row
        for i, label in enumerate(headers, start=1):
            if label not in block.measure_fields:
                continue
            self._put(FIRST_COL + i, totals.get(label), note_f, block.number_format(label))
        self.row += 1

        self._put(FIRST_COL, "Check", ctrl_f)
        for i, label in enumerate(headers, start=1):
            if label not in block.measure_fields:
                continue
            letter = col_letter(FIRST_COL + i)
            self._formula(
                FIRST_COL + i,
                f'=IF(ABS({letter}{control_row}-{letter}{expected_row})<=0.5,"OK","MISMATCH")',
                "OK", None, ctrl_f,
            )
        self.row += 1

    def _write_excluded_note(self, block: Block):
        if not block.excluded_records:
            return
        totals = {
            m: sum(r.values[m] for r in block.excluded_records
                   if isinstance(r.values.get(m), (int, float)))
            for m in block.measure_fields
        }
        self._put(FIRST_COL, "Excluded records, measure totals", note_f)
        for i, label in enumerate(block.fields, start=1):
            if label not in block.measure_fields:
                continue
            self._put(FIRST_COL + i, totals.get(label), note_f, block.number_format(label))
        self.row += 1
        self._line(
            "Shown as an independent control: an excluded total row should equal the "
            "extracted total above. No assumption is made about what the excluded records are."
        )

    # ────────────────────────────────────────────────────────── step 2
    def write_step2(self, result: Step2Result) -> None:
        block, ds = result.block, result.block.dataset
        self.row += GAP
        self._put(FIRST_COL, anchor_tag(ds.key, "STEP2"), anchor_f)
        self.row += 1
        self._put(FIRST_COL, "STEP 2 — NORMALISED", title_f, fill=step2_fill)
        self.row += 1

        for note in result.notes:
            self._line(f"· {note}")
        if result.assumed_fields:
            self._line(
                f"Confidence: {result.confidence.value} — step 1 read "
                f"{block.confidence.value}, but {len(result.assumed_fields)} column(s) "
                f"rest on a declared ratio rather than on a reported figure"
            )
        else:
            self._line(f"Confidence carried from step 1: {block.confidence.value}")
        self.row += 1

        columns = [block.provenance_label, *result.columns]
        for i, label in enumerate(columns):
            self._put(FIRST_COL + i, label, head_f, fill=step2_fill)
        self.row += 1

        calc_names = ([c.name for c in result.spec.calculations]
                      + [c.name for c in result.spec.cumulative])
        col_of = {label: FIRST_COL + 1 + i for i, label in enumerate(result.columns)}

        first_data = self.row
        for n, record in enumerate(result.records):
            self._put(FIRST_COL, record.source_ref, body_f)
            for label in result.declared_columns:
                self._put(col_of[label], record.values.get(label), body_f,
                          block.number_format(label))
            for calc in result.spec.calculations:
                expected = result.computed[calc.name][n]
                formula = calc.expression
                for fld in result.declared_columns:
                    formula = formula.replace(
                        "{" + fld + "}", f"{col_letter(col_of[fld])}{self.row}"
                    )
                self._formula(col_of[calc.name], f"=IFERROR({formula},\"\")",
                              expected, calc.number_format, body_f)
            self.row += 1
        last_data = self.row - 1

        # Cumulative shares, written as live ranges so a reviewer can see the running
        # sum and the total it is divided by — spec §9.2.1 S17.
        for cum in result.spec.cumulative:
            letter = col_letter(col_of[cum.field])
            for n, _ in enumerate(result.records):
                row = first_data + n
                self.row = row
                self._formula(
                    col_of[cum.name],
                    f"=IFERROR(SUM({letter}${first_data}:{letter}{row})"
                    f"/SUM({letter}${first_data}:{letter}${last_data}),\"\")",
                    result.computed[cum.name][n], cum.number_format, body_f,
                )
            self.row = last_data + 1

        measures = list(block.measure_fields)
        self._put(FIRST_COL, f"Control  (n = {len(result.records)})", ctrl_f, fill=ctrl_fill)
        control_row = self.row
        totals = result.totals()
        for label in measures:
            letter = col_letter(col_of[label])
            self._formula(col_of[label], f"=SUM({letter}{first_data}:{letter}{last_data})",
                          totals.get(label), block.number_format(label), ctrl_f, ctrl_fill)
        self.row += 1

        self._put(FIRST_COL, "Tie-back to step 1 (must be unchanged)", note_f)
        for label in measures:
            self._put(col_of[label], block.totals().get(label), note_f,
                      block.number_format(label))
        self.row += 1

        self._put(FIRST_COL, "Check", ctrl_f)
        tie_row = self.row - 1
        for label in measures:
            letter = col_letter(col_of[label])
            self._formula(
                col_of[label],
                f'=IF(ABS({letter}{control_row}-{letter}{tie_row})<=0.5,"OK","MISMATCH")',
                "OK", None, ctrl_f,
            )
        self.row += 1
        if calc_names:
            self._line(
                f"{', '.join(calc_names)} "
                f"{'is' if len(calc_names) == 1 else 'are'} value-adding, so no "
                "tie-back applies; each is written as a live formula over the "
                "columns above."
            )
        self._write_aggregate(result)
        self._write_figures(result)
        self._write_crosschecks(result.block)

    def _write_aggregate(self, result: Step2Result) -> None:
        """Further step-2 tables, one per declared aggregate — spec §9.2.1 S14."""
        for table in result.aggregates:
            self._write_one_aggregate(result, table)

    def _write_one_aggregate(self, result: Step2Result, table) -> None:
        self.row += 2
        self._put(FIRST_COL, table.title, head_f, fill=step2_fill)
        self.row += 1
        if table.note:
            self._line(table.note)
        if table.zero_filled:
            self._line(f"Years with no record are shown as 0: {', '.join(table.zero_filled)}")

        self._put(FIRST_COL, table.group_by, head_f, fill=step2_fill)
        for i, measure in enumerate(table.measures, start=1):
            self._put(FIRST_COL + i, measure, head_f, fill=step2_fill)
        self.row += 1

        first = self.row
        for key, values in table.rows:
            self._put(FIRST_COL, key, body_f, "@")
            for i, measure in enumerate(table.measures, start=1):
                self._put(FIRST_COL + i, values.get(measure, 0.0), body_f,
                          result.block.number_format(measure))
            self.row += 1
        last = self.row - 1

        self._put(FIRST_COL, f"Control  (n = {len(table.rows)})", ctrl_f, fill=ctrl_fill)
        totals = table.totals()
        for i, measure in enumerate(table.measures, start=1):
            letter = col_letter(FIRST_COL + i)
            self._formula(FIRST_COL + i, f"=SUM({letter}{first}:{letter}{last})",
                          totals.get(measure), result.block.number_format(measure),
                          ctrl_f, ctrl_fill)
        self.row += 1
        self._line("Grouping is value-preserving: this total must equal the detail total above.")

    def _write_figures(self, result: Step2Result) -> None:
        """Block-level derived figures — spec §9.2.1."""
        if not result.figures:
            return
        self.row += 1
        self._put(FIRST_COL, "Derived figures", head_f, fill=step2_fill)
        self.row += 1
        for figure in result.figures:
            self._put(FIRST_COL, figure.name, ctrl_f)
            if figure.value is None:
                self._put(FIRST_COL + 1, "not computed", note_f)
            else:
                self._put(FIRST_COL + 1, figure.value, ctrl_f, figure.number_format)
            self._put(FIRST_COL + 2, figure.detail, note_f)
            self._put(FIRST_COL + 4, figure.note, note_f)
            self.row += 1

    def _write_crosschecks(self, block: Block) -> None:
        """Interdependencies this sheet takes part in — spec §10.1."""
        if not block.crosschecks:
            return
        self.row += 1
        self._put(FIRST_COL, "Crosschecks against other sheets", head_f, fill=step2_fill)
        self.row += 1
        for result in block.crosschecks:
            self._put(FIRST_COL, result.label, ctrl_f)
            self._put(FIRST_COL + 1, result.status.upper(),
                      ctrl_f, fill=None if result.status == "passed" else ctrl_fill)
            self._put(FIRST_COL + 2,
                      f"{result.rule.left} {result.rule.relation} {result.rule.right}", note_f)
            self._put(FIRST_COL + 5, result.detail, note_f)
            self.row += 1


def write_blocks(ws, block: Block, result: Step2Result,
                 clear: bool = True) -> list[tuple[str, str, float]]:
    """Write step 1 and step 2 — spec §9.1.

    Both steps or neither. A dataset with no step-2 spec is not written at all and the
    run reports an error: a half-transformed sheet in an output workbook invites the
    reader to treat it as finished, and nothing in the sheet itself would say otherwise.

    ``clear`` removes output from an earlier run. A sheet carrying several
    section-blocks is written one block at a time, so only the first call may clear —
    otherwise each block would erase the block written before it.
    """
    from .extract import last_non_empty_row

    if clear:
        clear_generated(ws)
    start = last_non_empty_row(ws) + 1 + GAP
    writer = BlockWriter(ws, start)
    writer.write_step1(block)
    writer.write_step2(result)

    _widen(ws)
    return writer.formula_values


def _widen(ws) -> None:
    for col in range(FIRST_COL, FIRST_COL + 8):
        letter = col_letter(col)
        if ws.column_dimensions[letter].width in (None, 0):
            ws.column_dimensions[letter].width = 16


def write_loss_share(ws, table) -> list[tuple[str, str, float]]:
    """The section group's loss check, at the end of sheet 01 — spec §10.3.

    Written last and separated by the usual three blank rows, because it is about the
    treaty rather than about one block: it sums every loss dataset of the group and
    holds the total against the incurred losses the same group reports.
    """
    from .extract import last_non_empty_row
    from .lossshare import EXCEEDS, OK, ROLE_LABEL, TOLERANCE, WARNING

    writer = BlockWriter(ws, last_non_empty_row(ws) + 1 + GAP)
    writer._put(FIRST_COL, anchor_tag(f"{table.kind} sections", "LOSSSHARE"), anchor_f)
    writer.row += 1
    writer._put(FIRST_COL, table.title.upper(), title_f, fill=ctrl_fill)
    writer.row += 1

    writer._line(f"Sections summed: {', '.join(table.sections)}"
                 + (f" · scale {table.scale}" if table.scale else ""))
    if len(table.sheets) > 1:
        # The group spans sheets, so the same table is written at the end of each of
        # them. Saying so stops a reader treating the second copy as a second finding.
        others = [s for s in table.sheets if s != ws.title]
        writer._line("The same table is written at the end of "
                     + ", ".join(others) + " — one check, one result, shown wherever "
                     "a section of this group is reported.")
    if table.skipped:
        writer._line(f"NOT EVALUATED — {table.skipped}. Summing figures on a differing "
                     "basis would produce a number that means nothing.")
        _widen(ws)
        return writer.formula_values

    writer._line(
        f"Declared losses = {' + '.join(ROLE_LABEL[r] for r in table.roles)}; "
        f"total incurred = Incurred Losses in 01 for the same year."
    )
    writer._line(
        f"Two questions per year: declared ≤ incurred (tolerance {TOLERANCE:,.0f}), and "
        f"declared ÷ incurred ≤ {table.threshold:.0%} — above that the year is driven by "
        "single events rather than attrition, which changes how it is rated."
    )

    # The per-role columns show how the declared total is made up. With one role they
    # would simply repeat it, so they are written only where there is something to split.
    split = table.roles if len(table.roles) > 1 else ()

    headers = ["Year"] + [ROLE_LABEL[r] for r in split] + \
              ["Declared losses", "Total incurred (01)", "Share", "Status"]
    for i, text in enumerate(headers):
        writer._put(FIRST_COL + i, text, head_f, fill=ctrl_fill)
    writer.row += 1

    share_col = FIRST_COL + len(split) + 3
    status_col = share_col + 1
    first = writer.row
    for row in table.rows:
        writer._put(FIRST_COL, str(row.year), body_f, "@")
        for i, role in enumerate(split, start=1):
            writer._put(FIRST_COL + i, row.by_role.get(role, 0.0), body_f, "#,##0")
        declared_col = FIRST_COL + len(split) + 1
        writer._put(declared_col, row.declared, ctrl_f, "#,##0")
        writer._put(declared_col + 1, row.incurred, body_f, "#,##0")
        if row.share is not None:
            writer._formula(
                share_col,
                f"={col_letter(declared_col)}{writer.row}/"
                f"{col_letter(declared_col + 1)}{writer.row}",
                row.share, "0.0%", ctrl_f,
            )
        else:
            writer._put(share_col, "n/a", note_f)
        status = row.status(table.threshold)
        writer._put(status_col, status, ctrl_f,
                    fill=None if status == OK else ctrl_fill)
        writer.row += 1
    last = writer.row - 1

    writer._put(FIRST_COL, f"Control  (n = {len(table.rows)})", ctrl_f, fill=ctrl_fill)
    for i in range(1, len(split) + 3):
        letter = col_letter(FIRST_COL + i)
        writer._formula(FIRST_COL + i, f"=SUM({letter}{first}:{letter}{last})",
                        _column_total(table, split, i), "#,##0", ctrl_f, ctrl_fill)
    writer.row += 1

    partial = [str(r.year) for r in table.rows if r.partial]
    if partial:
        writer._line(
            f"Only some sections of this group report a history row for "
            f"{', '.join(partial)}, so the incurred total for those years covers fewer "
            "sections than the declared losses do. The comparison still runs; it is "
            "conservative, and the gap is named rather than closed by assumption."
        )

    grouped = {}
    for row in table.rows:
        status = row.status(table.threshold)
        if status != OK:
            grouped.setdefault(status, []).append(str(row.year))
    writer._line(
        "Every year within both limits."
        if not grouped else
        "; ".join(f"{status}: {', '.join(years)}" for status, years in grouped.items())
        + ". "
        + ("A declared total above the incurred total cannot be right — one of the two "
           "sheets is wrong, or they are on different bases."
           if table.worst == EXCEEDS else
           "Not an error: a fact about the portfolio, raised so it is priced knowingly."
           if table.worst == WARNING else
           "A year in 01 has no incurred figure to compare against.")
    )
    _widen(ws)
    return writer.formula_values


def _column_total(table, split, offset: int):
    """The tool's own value for a written SUM formula, so the cache stays honest."""
    if offset <= len(split):
        role = split[offset - 1]
        return sum(r.by_role.get(role, 0.0) for r in table.rows)
    declared, incurred = table.totals()
    return declared if offset == len(split) + 1 else incurred


def write_growth(ws, table) -> list[tuple[str, str, float]]:
    """Exposure against premium, at the end of the aggregate sheet — spec §10.4.

    Growth is never an error here. A portfolio that shrinks is a fact about the book, so
    the block reports and warns; nothing in it can fail a run.
    """
    from .extract import last_non_empty_row
    from .growth import NO_BASIS, OK

    writer = BlockWriter(ws, last_non_empty_row(ws) + 1 + GAP)
    writer._put(FIRST_COL, anchor_tag(f"{table.section} {table.role}", "GROWTH"), anchor_f)
    writer.row += 1
    writer._put(FIRST_COL, table.title.upper(), title_f, fill=ctrl_fill)
    writer.row += 1

    if table.skipped:
        writer._line(f"NOT EVALUATED — {table.skipped}.")
        _widen(ws)
        return writer.formula_values

    writer._line(
        "Exposure is the sum of the zone totals of each version; premium is the "
        f"expiring year from 01 and the renewal year from 02, since 01 carries no "
        f"forward figure."
    )
    writer._line(
        f"Implied rate change = (1 + premium growth) ÷ (1 + exposure growth) − 1. "
        f"Beyond ±{table.threshold:.0%} it is flagged — a warning, never an error: a book "
        "may shrink or grow for good reasons."
    )

    for i, text in enumerate(["Version", "As at", "Exposure", "Change"]):
        writer._put(FIRST_COL + i, text, head_f, fill=ctrl_fill)
    writer.row += 1
    for version in table.versions:
        writer._put(FIRST_COL, version.period, body_f, "@")
        writer._put(FIRST_COL + 1, version.as_at, body_f, "@")
        writer._put(FIRST_COL + 2, version.exposure, body_f, "#,##0")
        if version.change is not None:
            writer._put(FIRST_COL + 3, version.change, body_f, "0.0%")
        writer.row += 1

    writer.row += 1
    for label, pair in (("Premium, expiring year", table.premium_from),
                        ("Premium, renewal year", table.premium_to)):
        writer._put(FIRST_COL, label, note_f)
        if pair is not None:
            writer._put(FIRST_COL + 1, pair[0], note_f, "@")
            writer._put(FIRST_COL + 2, pair[1], note_f, "#,##0")
        else:
            writer._put(FIRST_COL + 1, "not found", note_f)
        writer.row += 1

    writer.row += 1
    for label, value, fmt in (
        ("Exposure growth", table.exposure_growth, "0.0%"),
        ("Premium growth", table.premium_growth, "0.0%"),
        ("Implied rate change", table.implied_rate_change, "0.0%"),
    ):
        writer._put(FIRST_COL, label, ctrl_f)
        if value is None:
            writer._put(FIRST_COL + 2, "not computed", note_f)
        else:
            writer._put(FIRST_COL + 2, value, ctrl_f, fmt)
        writer.row += 1

    writer._put(FIRST_COL, "Status", ctrl_f)
    writer._put(FIRST_COL + 2, table.status, ctrl_f,
                fill=None if table.status == OK else ctrl_fill)
    writer.row += 1

    if table.status == OK:
        writer._line("Premium and exposure have moved together within the declared band.")
    elif table.status == NO_BASIS:
        writer._line("One side is missing, so no rate movement can be read.")
    else:
        rate = table.implied_rate_change
        direction = "fallen" if rate < 0 else "risen"
        writer._line(
            f"Premium per unit of exposure has {direction} {abs(rate):.1%}. Not an error "
            "— but worth holding against the cedent's own rate change in 09 before the "
            "renewal is priced."
        )
    _widen(ws)
    return writer.formula_values
