"""Orchestration and logging. Specification_v1.md §9, §11, §12."""

from __future__ import annotations

import hashlib
import logging
import shutil
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

from .crosschecks import hypotheses_from, results_for, run_rules
from .extract import _number_hypotheses, extract_sheet, last_non_empty_row
from .markers import read_markers
from .model import Block, ExtractionError
from .nomenclature import Nomenclature, read_nomenclature, sheet_sort_key
from .recalc import inject
from .specs import SPEC_VERSION, step2_for
from .transform import Step2Result, apply_step2
from .writer import write_blocks

TOOL_VERSION = "0.1.0"


@dataclass
class SheetOutcome:
    sheet: str
    status: str                       # processed | skipped | error
    detail: str = ""
    blocks: list[Block] = field(default_factory=list)
    results: list[Step2Result] = field(default_factory=list)


@dataclass
class RunReport:
    source: Path
    output: Path
    source_sha256: str
    started: datetime
    outcomes: list[SheetOutcome] = field(default_factory=list)
    rule_results: list = field(default_factory=list)
    injected: int = 0

    @property
    def ok(self) -> bool:
        return all(o.status != "error" for o in self.outcomes)

    @property
    def rules_ok(self) -> bool:
        return all(r.status != "failed" for r in self.rule_results)


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(65536), b""):
            h.update(chunk)
    return h.hexdigest()


def setup_logs(log_dir: Path, stamp: str) -> tuple[logging.Logger, logging.Logger]:
    """Two logs, linked to the sheets by block ID — spec §11."""
    log_dir.mkdir(parents=True, exist_ok=True)

    debug = logging.getLogger("datatransform.debug")
    process = logging.getLogger("datatransform.process")
    for lg in (debug, process):
        lg.handlers.clear()
        lg.setLevel(logging.DEBUG)
        lg.propagate = False

    dh = logging.FileHandler(log_dir / f"{stamp}_debug.log", encoding="utf-8")
    dh.setFormatter(logging.Formatter("%(asctime)s  %(levelname)-7s  %(message)s"))
    debug.addHandler(dh)

    ph = logging.FileHandler(log_dir / f"{stamp}_process.log", encoding="utf-8")
    ph.setFormatter(logging.Formatter("%(message)s"))
    process.addHandler(ph)
    return debug, process


def preflight(values_wb, formulas_wb, debug) -> list[str]:
    """External links, error cells and merged ranges — spec §12 T3."""
    findings = []
    ext = getattr(formulas_wb, "_external_links", None) or []
    if ext:
        findings.append(f"workbook carries {len(ext)} external link(s)")
    for ws in formulas_wb.worksheets:
        if ws.merged_cells.ranges:
            findings.append(f"{ws.title}: {len(ws.merged_cells.ranges)} merged range(s)")
    for ws in values_wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.strip().startswith("#") \
                        and cell.value.strip().endswith(("!", "?", "A", "0")):
                    findings.append(f"{ws.title}!{cell.coordinate}: {cell.value}")
    for f in findings:
        debug.warning("pre-flight: %s", f)
    return findings


def run(source: str | Path, output: str | Path | None = None,
        log_dir: str | Path = "logs") -> RunReport:
    source = Path(source)
    started = datetime.now(timezone.utc)
    stamp = started.strftime("%Y%m%dT%H%M%SZ")
    output = Path(output) if output else source.with_name(f"{source.stem}_transformed.xlsx")

    debug, process = setup_logs(Path(log_dir), stamp)
    digest = sha256(source)

    process.info("=" * 78)
    process.info("Data_Transform run %s", stamp)
    process.info("source      : %s", source)
    process.info("sha256      : %s", digest)
    process.info("output      : %s", output)
    process.info("spec version: %s   tool version: %s", SPEC_VERSION, TOOL_VERSION)
    process.info("=" * 78)
    debug.info("run %s started; source=%s sha256=%s", stamp, source, digest)

    report = RunReport(source, output, digest, started)

    # Spec §9.1 O1 — the source workbook is never modified.
    if output.resolve() == source.resolve():
        raise ExtractionError(
            f"output would overwrite the source workbook ({source}); the source is the "
            "audit baseline and is never modified — choose a different output path"
        )
    output.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, output)

    values_wb = load_workbook(source, data_only=True)
    formulas_wb = load_workbook(source, data_only=False)
    out_wb = load_workbook(output, data_only=False)

    for finding in preflight(values_wb, formulas_wb, debug):
        process.info("pre-flight  : %s", finding)

    nomenclature = read_nomenclature(formulas_wb)
    debug.info("nomenclature: %d dataset(s), %d vocabulary entries",
               len(nomenclature.datasets), len(nomenclature.vocabulary))
    process.info("")
    process.info("Sheet 00 declares %d dataset(s): %s",
                 len(nomenclature.datasets), ", ".join(sorted(nomenclature.datasets)))

    formula_values: list[tuple[str, str, float]] = []

    # Pass 1 — extract every sheet, so the rules have both sides available.
    for title in sorted(out_wb.sheetnames, key=sheet_sort_key):
        if sheet_sort_key(title)[0] == "000":
            continue
        report.outcomes.append(
            _extract_sheet_outcome(title, values_wb, formulas_wb, nomenclature, debug, process)
        )

    blocks = [b for o in report.outcomes for b in o.blocks]
    report.rule_results = run_rules(nomenclature, blocks)
    _log_rules(report.rule_results, nomenclature, debug, process)

    for outcome in report.outcomes:
        for block in outcome.blocks:
            block.crosschecks = results_for(report.rule_results, block)
            block.hypotheses.extend(hypotheses_from(report.rule_results, block))
            _number_hypotheses(block)

    # Pass 2 — transform and write.
    for outcome in report.outcomes:
        _write_sheet(outcome, out_wb, nomenclature, blocks,
                     debug, process, formula_values)

    out_wb.save(output)
    result = inject(output, formula_values)
    report.injected = result["injected"]
    debug.info("cached %d formula value(s); unresolved=%s",
               result["injected"], result["unresolved"] or "none")

    process.info("")
    process.info("-" * 78)
    for o in report.outcomes:
        process.info("%-26s %-10s %s", o.sheet, o.status, o.detail)
    process.info("run %s", "completed" if report.ok else "completed WITH ERRORS")
    return report


def _extract_sheet_outcome(title, values_wb, formulas_wb, nomenclature,
                           debug, process) -> SheetOutcome:
    dataset = nomenclature.dataset_for(title)
    if dataset is None:
        debug.info("%s: not declared in sheet 00 — skipped", title)
        return SheetOutcome(title, "skipped", "not declared in sheet 00")

    values_ws = values_wb[title]
    markers = read_markers(formulas_wb[title], last_non_empty_row(values_ws))
    if not any(m.name == "Header" for m in markers):
        debug.info("%s: no Header_i marker — nothing extracted", title)
        process.info("")
        process.info("%s — no Header_i in column A, so nothing is extracted (spec §4 M7).", title)
        return SheetOutcome(title, "skipped", "no Header_i marker")

    try:
        blocks = extract_sheet(values_ws, formulas_wb[title], nomenclature, markers)
    except ExtractionError as exc:
        debug.error("%s: %s", title, exc)
        process.info("")
        process.info("%s — FAILED: %s", title, exc)
        return SheetOutcome(title, "error", str(exc))

    outcome = SheetOutcome(title, "processed", blocks=list(blocks))
    outcome.detail = (
        f"{len(outcome.blocks)} block(s), "
        f"{sum(len(b.records) for b in outcome.blocks)} record(s)"
    )
    return outcome


def _write_sheet(outcome, out_wb, nomenclature, blocks,
                 debug, process, formula_values) -> None:
    for block in outcome.blocks:
        spec = step2_for(block.dataset.key)
        if spec is None:
            outcome.status = "error"
            outcome.detail = f"no step-2 spec for dataset {block.dataset.key!r}"
            process.info("%s — no step-2 spec for %s", outcome.sheet, block.dataset.key)
            continue

        result = apply_step2(block, spec, nomenclature, blocks)
        formula_values.extend(write_blocks(out_wb[outcome.sheet], block, result))
        outcome.results.append(result)
        _log_block(block, result, debug, process)


def _log_rules(results, nomenclature, debug, process) -> None:
    if not results:
        return
    process.info("")
    sections = ", ".join(f"{s.name} ({s.kind})" for s in nomenclature.sections) or "none declared"
    process.info("Crosschecks (⟦RULES⟧ of sheet 00, N = %s) — sections: %s",
                 nomenclature.actual_year, sections)
    for r in results:
        process.info("  %-11s %-15s %s %s %s — %s",
                     r.label, r.status.upper(), r.rule.left, r.rule.relation,
                     r.rule.right, r.detail)
        debug.info("rule %s: %s (%s)", r.label, r.status, r.detail)


def _log_block(block: Block, result: Step2Result, debug, process):
    key = f"{block.dataset.key}/block {block.index}"
    debug.info("%s: orientation=%s header=%s info=%s map=%s",
               key, block.orientation.value, block.header_ref, block.info_ref, block.address_map)
    debug.info("%s: candidates=%d extracted=%d excluded=%d unextracted=%s",
               key, block.candidates, len(block.records), block.excluded,
               block.unextracted or "none")
    for r in block.records:
        debug.debug("%s: %s -> %s", key, r.source_ref, r.values)

    process.info("")
    process.info("%s — STEP 1", key)
    process.info("  %s block; extraction %s %s; selector %s.",
                 block.orientation.value,
                 "row" if block.orientation.value == "row-wise" else "column",
                 block.header_ref, block.info_ref)
    process.info("  Fields resolved by label: %s.",
                 ", ".join(f"{k} → {v}" for k, v in block.address_map.items()))
    process.info("  %d candidate record(s); %d extracted, %d excluded because the selector "
                 "was empty.", block.candidates, len(block.records), block.excluded)
    process.info("  Types applied: %s.",
                 ", ".join(f"{h} read as {block.field_types[h].value}"
                           for h in block.dataset.headers))
    for c in block.coercions:
        debug.debug("%s: converted %s at %s: %r -> %r (%s)",
                    key, c.field, c.source_ref, c.before, c.after, c.note)
        if not c.routine:
            process.info("  Converted %s at %s: %r → %r — %s",
                         c.field, c.source_ref, c.before, c.after, c.note)
    if block.unextracted:
        process.info("  Contained data but was not extracted (not declared in sheet 00): %s. "
                     "Should it be?", ", ".join(block.unextracted))
    for a in block.attributes.values():
        process.info("  Attribute %s = %s%s", a.name, a.value,
                     " — declared as a hypothesis" if a.is_hypothesis else "")
    for h in block.hypotheses:
        process.info("  %s  %s = %s  [%s, %s] — %s",
                     h.id, h.attribute, h.value, h.confidence.value, h.status, h.note)
    process.info("  Dataset confidence: %s (worst of all hypotheses).", block.confidence.value)
    for measure, total in block.totals().items():
        process.info("  Control %s: %s", measure, f"{total:,.0f}")

    process.info("")
    process.info("%s — STEP 2", key)
    for note in result.notes:
        process.info("  %s", note)
    for measure, total in result.totals().items():
        process.info("  Control %s: %s (unchanged from step 1)", measure, f"{total:,.0f}")
